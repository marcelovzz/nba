import io
import math
from typing import Any, Dict, List, Optional, Tuple
import numpy as np
import pandas as pd

try:
    import streamlit as st
except ImportError:
    st = None
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

LEAGUE_AVG_ORTG = 115.5
LEAGUE_AVG_PACE = 99.0
LEAGUE_AVG_ISO_PPP = 1.03
LEAGUE_AVG_PNR_PPP = 1.04
LEAGUE_AVG_TOV = 13.1

# Desviaciones estándar empíricas de simulaciones Monte Carlo NBA (5,000 posesiones)
NBA_TOTAL_SIM_STD = 13.6
NBA_MARGIN_SIM_STD = 11.8

# Límites de seguridad (clipping) para ratings de la NBA moderna en Calibración Kalman
PACE_CLIP_RANGE: Tuple[float, float] = (92.0, 106.5)
ORTG_CLIP_RANGE: Tuple[float, float] = (104.0, 126.5)
DRTG_CLIP_RANGE: Tuple[float, float] = (104.0, 124.5)

GAME_SLATE_TIMES: List[str] = [
    "19:00 ET", "19:30 ET", "19:30 ET", "20:00 ET", "20:00 ET",
    "20:30 ET", "20:30 ET", "21:00 ET", "21:00 ET", "21:30 ET",
    "22:00 ET", "22:00 ET", "22:30 ET", "22:30 ET", "23:00 ET",
]

LINEUP_COLS: List[str] = ["Posición", "Jugador", "Minutos_Esperados", "O_EPM", "D_EPM"]
EXCEL_TEMPLATE_BASE_COLS: List[str] = [
    "Team_Abbr", "Pos", "Jugador", "Minutos", "O_EPM", "D_EPM", "Pace_Diff", "Es_Estrella", "PPG_Linea",
]
EXCEL_TEMPLATE_COLS: List[str] = [
    "Team_Abbr", "Pos", "Jugador", "Minutos", "Minutos_Base",
    "O_EPM", "D_EPM", "Pace_Diff", "Es_Estrella",
    "PPG_Base", "RPG_Base", "APG_Base", "PPG_Linea",
]
ROSTER_EDITOR_COLS: List[str] = [
    "Team_Abbr", "Star1_Name", "Star1_PPG", "Star2_Name", "Star2_PPG", "Pace", "Off_Rtg", "Def_Rtg",
]

FULL_30_TEAM_SLATE: List[Tuple[str, str]] = [
    ("Detroit Pistons", "Boston Celtics"), ("New York Knicks", "Miami Heat"),
    ("Dallas Mavericks", "Denver Nuggets"), ("Golden State Warriors", "Los Angeles Lakers"),
    ("Minnesota Timberwolves", "Oklahoma City Thunder"), ("Philadelphia 76ers", "Cleveland Cavaliers"),
    ("Milwaukee Bucks", "Indiana Pacers"), ("Phoenix Suns", "LA Clippers"),
    ("Sacramento Kings", "Memphis Grizzlies"), ("Houston Rockets", "San Antonio Spurs"),
    ("Orlando Magic", "Atlanta Hawks"), ("New Orleans Pelicans", "Chicago Bulls"),
    ("Brooklyn Nets", "Toronto Raptors"), ("Charlotte Hornets", "Washington Wizards"),
    ("Portland Trail Blazers", "Utah Jazz"),
]

# ==============================================================================
# MÉTRICAS BASE DE REBOTES (RPG_Base) Y ASISTENCIAS (APG_Base) PARA ESTRELLAS NBA
# ==============================================================================
STAR_BASE_REB_AST_MAP: Dict[str, Tuple[float, float]] = {
    "jayson tatum": (8.6, 4.9),
    "jaylen brown": (5.8, 3.7),
    "cam thomas": (3.4, 3.1),
    "dennis schroder": (3.1, 6.4),
    "jalen brunson": (3.6, 6.9),
    "karl-anthony towns": (11.2, 3.1),
    "joel embiid": (11.0, 5.4),
    "tyrese maxey": (3.8, 6.3),
    "scottie barnes": (8.2, 6.1),
    "rj barrett": (6.2, 4.1),
    "zach lavine": (5.1, 4.2),
    "coby white": (4.6, 5.2),
    "donovan mitchell": (5.1, 5.9),
    "darius garland": (2.9, 6.6),
    "cade cunningham": (6.4, 8.2),
    "jaden ivey": (3.9, 4.2),
    "tyrese haliburton": (3.9, 10.4),
    "pascal siakam": (7.3, 4.1),
    "giannis antetokounmpo": (11.6, 6.4),
    "damian lillard": (4.4, 7.1),
    "nikola jokic": (12.4, 9.2),
    "jamal murray": (4.1, 6.5),
    "anthony edwards": (5.6, 5.1),
    "julius randle": (9.0, 4.7),
    "shai gilgeous-alexander": (5.5, 6.3),
    "jalen williams": (5.2, 4.8),
    "anfernee simons": (3.4, 5.4),
    "jerami grant": (3.8, 2.7),
    "lauri markkanen": (8.2, 2.1),
    "keyonte george": (3.2, 5.6),
    "stephen curry": (4.5, 5.4),
    "jonathan kuminga": (5.2, 2.3),
    "james harden": (5.8, 8.6),
    "norman powell": (3.2, 2.2),
    "lebron james": (7.4, 8.2),
    "anthony davis": (12.5, 3.6),
    "kevin durant": (6.6, 5.1),
    "devin booker": (4.6, 6.8),
    "de'aaron fox": (4.8, 5.8),
    "demar derozan": (4.4, 5.1),
    "trae young": (3.1, 10.8),
    "jalen johnson": (9.4, 5.1),
    "lamelo ball": (5.3, 8.0),
    "brandon miller": (4.5, 3.1),
    "jimmy butler": (5.6, 5.0),
    "bam adebayo": (10.4, 4.1),
    "paolo banchero": (7.2, 5.4),
    "franz wagner": (5.7, 4.4),
    "jordan poole": (2.9, 4.6),
    "kyle kuzma": (6.6, 4.1),
    "luka doncic": (9.1, 9.6),
    "kyrie irving": (5.0, 5.3),
    "alperen sengun": (9.6, 5.0),
    "jalen green": (5.1, 3.6),
    "ja morant": (5.6, 8.2),
    "jaren jackson jr.": (6.1, 2.3),
    "zion williamson": (6.2, 5.0),
    "brandon ingram": (5.2, 5.6),
    "victor wembanyama": (10.8, 3.9),
    "devin vassell": (3.9, 4.1),
}

# ==============================================================================
# BASE DE DATOS PREDETERMINADA: 30 FRANQUICIAS + QUINTETOS TITULARES EPM
# ==============================================================================
DEFAULT_CSV_DATA: List[Dict] = [
    {"Team": "Boston Celtics", "Team_Abbr": "BOS", "Conference": "East", "Division": "Atlántico", "Pace": 98.4, "Off_Rtg": 122.2, "Def_Rtg": 110.6, "3PA_Rate": 0.472, "3P_Pct": 0.388, "TOV_Pct": 11.9, "Iso_PPP_Off": 1.12, "Iso_PPP_Def": 0.91, "PnR_PPP_Off": 1.09, "PnR_PPP_Def": 0.93, "Star1_Name": "Jayson Tatum", "Star1_PPG": 27.1, "Star1_Line": 26.5, "Star1_Style": "ISO", "Star1_Std": 5.9, "Star2_Name": "Jaylen Brown", "Star2_PPG": 23.2, "Star2_Line": 22.5, "Star2_Style": "ISO", "Star2_Std": 5.4},
    {"Team": "Brooklyn Nets", "Team_Abbr": "BKN", "Conference": "East", "Division": "Atlántico", "Pace": 97.6, "Off_Rtg": 112.1, "Def_Rtg": 116.4, "3PA_Rate": 0.418, "3P_Pct": 0.362, "TOV_Pct": 13.8, "Iso_PPP_Off": 1.01, "Iso_PPP_Def": 1.07, "PnR_PPP_Off": 0.99, "PnR_PPP_Def": 1.08, "Star1_Name": "Cam Thomas", "Star1_PPG": 24.2, "Star1_Line": 23.5, "Star1_Style": "ISO", "Star1_Std": 6.3, "Star2_Name": "Dennis Schroder", "Star2_PPG": 17.6, "Star2_Line": 16.5, "Star2_Style": "PNR", "Star2_Std": 4.9},
    {"Team": "New York Knicks", "Team_Abbr": "NYK", "Conference": "East", "Division": "Atlántico", "Pace": 96.2, "Off_Rtg": 118.5, "Def_Rtg": 112.4, "3PA_Rate": 0.405, "3P_Pct": 0.376, "TOV_Pct": 12.6, "Iso_PPP_Off": 1.10, "Iso_PPP_Def": 0.96, "PnR_PPP_Off": 1.11, "PnR_PPP_Def": 0.95, "Star1_Name": "Jalen Brunson", "Star1_PPG": 28.7, "Star1_Line": 27.5, "Star1_Style": "PNR", "Star1_Std": 6.1, "Star2_Name": "Karl-Anthony Towns", "Star2_PPG": 22.1, "Star2_Line": 21.5, "Star2_Style": "PNR", "Star2_Std": 5.3},
    {"Team": "Philadelphia 76ers", "Team_Abbr": "PHI", "Conference": "East", "Division": "Atlántico", "Pace": 97.1, "Off_Rtg": 116.8, "Def_Rtg": 113.2, "3PA_Rate": 0.384, "3P_Pct": 0.369, "TOV_Pct": 12.2, "Iso_PPP_Off": 1.11, "Iso_PPP_Def": 0.99, "PnR_PPP_Off": 1.08, "PnR_PPP_Def": 0.98, "Star1_Name": "Joel Embiid", "Star1_PPG": 29.8, "Star1_Line": 28.5, "Star1_Style": "ISO", "Star1_Std": 6.5, "Star2_Name": "Tyrese Maxey", "Star2_PPG": 26.1, "Star2_Line": 25.5, "Star2_Style": "PNR", "Star2_Std": 5.8},
    {"Team": "Toronto Raptors", "Team_Abbr": "TOR", "Conference": "East", "Division": "Atlántico", "Pace": 100.4, "Off_Rtg": 112.8, "Def_Rtg": 117.1, "3PA_Rate": 0.366, "3P_Pct": 0.351, "TOV_Pct": 14.4, "Iso_PPP_Off": 0.97, "Iso_PPP_Def": 1.08, "PnR_PPP_Off": 1.01, "PnR_PPP_Def": 1.09, "Star1_Name": "Scottie Barnes", "Star1_PPG": 20.6, "Star1_Line": 19.5, "Star1_Style": "ISO", "Star1_Std": 5.2, "Star2_Name": "RJ Barrett", "Star2_PPG": 21.4, "Star2_Line": 20.5, "Star2_Style": "PNR", "Star2_Std": 5.4},
    {"Team": "Chicago Bulls", "Team_Abbr": "CHI", "Conference": "East", "Division": "Central", "Pace": 101.2, "Off_Rtg": 114.2, "Def_Rtg": 116.5, "3PA_Rate": 0.412, "3P_Pct": 0.367, "TOV_Pct": 13.5, "Iso_PPP_Off": 1.03, "Iso_PPP_Def": 1.06, "PnR_PPP_Off": 1.04, "PnR_PPP_Def": 1.07, "Star1_Name": "Zach LaVine", "Star1_PPG": 22.9, "Star1_Line": 22.5, "Star1_Style": "ISO", "Star1_Std": 5.8, "Star2_Name": "Coby White", "Star2_PPG": 19.8, "Star2_Line": 19.5, "Star2_Style": "PNR", "Star2_Std": 5.3},
    {"Team": "Cleveland Cavaliers", "Team_Abbr": "CLE", "Conference": "East", "Division": "Central", "Pace": 99.6, "Off_Rtg": 121.0, "Def_Rtg": 111.4, "3PA_Rate": 0.422, "3P_Pct": 0.386, "TOV_Pct": 12.1, "Iso_PPP_Off": 1.10, "Iso_PPP_Def": 0.94, "PnR_PPP_Off": 1.13, "PnR_PPP_Def": 0.94, "Star1_Name": "Donovan Mitchell", "Star1_PPG": 26.8, "Star1_Line": 26.5, "Star1_Style": "PNR", "Star1_Std": 6.0, "Star2_Name": "Darius Garland", "Star2_PPG": 20.4, "Star2_Line": 19.5, "Star2_Style": "PNR", "Star2_Std": 5.1},
    {"Team": "Detroit Pistons", "Team_Abbr": "DET", "Conference": "East", "Division": "Central", "Pace": 100.6, "Off_Rtg": 110.4, "Def_Rtg": 118.2, "3PA_Rate": 0.358, "3P_Pct": 0.348, "TOV_Pct": 15.2, "Iso_PPP_Off": 0.94, "Iso_PPP_Def": 1.13, "PnR_PPP_Off": 0.97, "PnR_PPP_Def": 1.11, "Star1_Name": "Cade Cunningham", "Star1_PPG": 22.8, "Star1_Line": 24.5, "Star1_Style": "PNR", "Star1_Std": 5.6, "Star2_Name": "Jaden Ivey", "Star2_PPG": 16.4, "Star2_Line": 16.5, "Star2_Style": "ISO", "Star2_Std": 4.8},
    {"Team": "Indiana Pacers", "Team_Abbr": "IND", "Conference": "East", "Division": "Central", "Pace": 102.4, "Off_Rtg": 119.8, "Def_Rtg": 116.8, "3PA_Rate": 0.395, "3P_Pct": 0.378, "TOV_Pct": 12.7, "Iso_PPP_Off": 1.06, "Iso_PPP_Def": 1.08, "PnR_PPP_Off": 1.14, "PnR_PPP_Def": 1.09, "Star1_Name": "Tyrese Haliburton", "Star1_PPG": 20.8, "Star1_Line": 20.5, "Star1_Style": "PNR", "Star1_Std": 5.7, "Star2_Name": "Pascal Siakam", "Star2_PPG": 21.6, "Star2_Line": 21.5, "Star2_Style": "ISO", "Star2_Std": 5.1},
    {"Team": "Milwaukee Bucks", "Team_Abbr": "MIL", "Conference": "East", "Division": "Central", "Pace": 99.8, "Off_Rtg": 117.6, "Def_Rtg": 114.5, "3PA_Rate": 0.402, "3P_Pct": 0.375, "TOV_Pct": 12.8, "Iso_PPP_Off": 1.11, "Iso_PPP_Def": 1.03, "PnR_PPP_Off": 1.12, "PnR_PPP_Def": 1.04, "Star1_Name": "Giannis Antetokounmpo", "Star1_PPG": 30.8, "Star1_Line": 30.5, "Star1_Style": "ISO", "Star1_Std": 6.1, "Star2_Name": "Damian Lillard", "Star2_PPG": 25.1, "Star2_Line": 24.5, "Star2_Style": "PNR", "Star2_Std": 6.0},
    {"Team": "Denver Nuggets", "Team_Abbr": "DEN", "Conference": "West", "Division": "Noroeste", "Pace": 97.4, "Off_Rtg": 119.4, "Def_Rtg": 112.3, "3PA_Rate": 0.342, "3P_Pct": 0.374, "TOV_Pct": 12.1, "Iso_PPP_Off": 1.06, "Iso_PPP_Def": 0.98, "PnR_PPP_Off": 1.14, "PnR_PPP_Def": 0.97, "Star1_Name": "Nikola Jokic", "Star1_PPG": 26.6, "Star1_Line": 26.5, "Star1_Style": "PNR", "Star1_Std": 5.2, "Star2_Name": "Jamal Murray", "Star2_PPG": 21.2, "Star2_Line": 20.5, "Star2_Style": "PNR", "Star2_Std": 5.7},
    {"Team": "Minnesota Timberwolves", "Team_Abbr": "MIN", "Conference": "West", "Division": "Noroeste", "Pace": 97.8, "Off_Rtg": 115.8, "Def_Rtg": 109.8, "3PA_Rate": 0.410, "3P_Pct": 0.382, "TOV_Pct": 13.9, "Iso_PPP_Off": 1.08, "Iso_PPP_Def": 0.90, "PnR_PPP_Off": 1.06, "PnR_PPP_Def": 0.91, "Star1_Name": "Anthony Edwards", "Star1_PPG": 27.6, "Star1_Line": 26.5, "Star1_Style": "ISO", "Star1_Std": 6.2, "Star2_Name": "Julius Randle", "Star2_PPG": 20.9, "Star2_Line": 20.5, "Star2_Style": "ISO", "Star2_Std": 5.2},
    {"Team": "Oklahoma City Thunder", "Team_Abbr": "OKC", "Conference": "West", "Division": "Noroeste", "Pace": 100.5, "Off_Rtg": 120.4, "Def_Rtg": 109.2, "3PA_Rate": 0.398, "3P_Pct": 0.389, "TOV_Pct": 11.6, "Iso_PPP_Off": 1.14, "Iso_PPP_Def": 0.89, "PnR_PPP_Off": 1.11, "PnR_PPP_Def": 0.90, "Star1_Name": "Shai Gilgeous-Alexander", "Star1_PPG": 30.2, "Star1_Line": 29.5, "Star1_Style": "ISO", "Star1_Std": 5.4, "Star2_Name": "Jalen Williams", "Star2_PPG": 21.1, "Star2_Line": 20.5, "Star2_Style": "PNR", "Star2_Std": 4.9},
    {"Team": "Portland Trail Blazers", "Team_Abbr": "POR", "Conference": "West", "Division": "Noroeste", "Pace": 99.1, "Off_Rtg": 109.6, "Def_Rtg": 117.4, "3PA_Rate": 0.372, "3P_Pct": 0.346, "TOV_Pct": 15.0, "Iso_PPP_Off": 0.95, "Iso_PPP_Def": 1.10, "PnR_PPP_Off": 0.98, "PnR_PPP_Def": 1.09, "Star1_Name": "Anfernee Simons", "Star1_PPG": 21.5, "Star1_Line": 21.5, "Star1_Style": "PNR", "Star1_Std": 5.9, "Star2_Name": "Jerami Grant", "Star2_PPG": 19.2, "Star2_Line": 18.5, "Star2_Style": "ISO", "Star2_Std": 5.1},
    {"Team": "Utah Jazz", "Team_Abbr": "UTA", "Conference": "West", "Division": "Noroeste", "Pace": 100.2, "Off_Rtg": 113.2, "Def_Rtg": 118.9, "3PA_Rate": 0.408, "3P_Pct": 0.356, "TOV_Pct": 15.4, "Iso_PPP_Off": 0.99, "Iso_PPP_Def": 1.11, "PnR_PPP_Off": 1.02, "PnR_PPP_Def": 1.12, "Star1_Name": "Lauri Markkanen", "Star1_PPG": 23.1, "Star1_Line": 22.5, "Star1_Style": "PNR", "Star1_Std": 5.5, "Star2_Name": "Keyonte George", "Star2_PPG": 16.9, "Star2_Line": 16.5, "Star2_Style": "PNR", "Star2_Std": 5.2},
    {"Team": "Golden State Warriors", "Team_Abbr": "GSW", "Conference": "West", "Division": "Pacífico", "Pace": 100.8, "Off_Rtg": 117.3, "Def_Rtg": 113.8, "3PA_Rate": 0.456, "3P_Pct": 0.380, "TOV_Pct": 14.3, "Iso_PPP_Off": 1.04, "Iso_PPP_Def": 1.00, "PnR_PPP_Off": 1.10, "PnR_PPP_Def": 1.01, "Star1_Name": "Stephen Curry", "Star1_PPG": 26.8, "Star1_Line": 26.5, "Star1_Style": "PNR", "Star1_Std": 6.4, "Star2_Name": "Jonathan Kuminga", "Star2_PPG": 16.8, "Star2_Line": 16.5, "Star2_Style": "ISO", "Star2_Std": 4.9},
    {"Team": "LA Clippers", "Team_Abbr": "LAC", "Conference": "West", "Division": "Pacífico", "Pace": 96.9, "Off_Rtg": 115.6, "Def_Rtg": 113.1, "3PA_Rate": 0.389, "3P_Pct": 0.376, "TOV_Pct": 13.2, "Iso_PPP_Off": 1.11, "Iso_PPP_Def": 0.99, "PnR_PPP_Off": 1.09, "PnR_PPP_Def": 0.98, "Star1_Name": "James Harden", "Star1_PPG": 21.6, "Star1_Line": 20.5, "Star1_Style": "ISO", "Star1_Std": 5.8, "Star2_Name": "Norman Powell", "Star2_PPG": 22.4, "Star2_Line": 21.5, "Star2_Style": "PNR", "Star2_Std": 5.4},
    {"Team": "Los Angeles Lakers", "Team_Abbr": "LAL", "Conference": "West", "Division": "Pacífico", "Pace": 101.4, "Off_Rtg": 115.9, "Def_Rtg": 115.2, "3PA_Rate": 0.349, "3P_Pct": 0.377, "TOV_Pct": 13.4, "Iso_PPP_Off": 1.07, "Iso_PPP_Def": 1.04, "PnR_PPP_Off": 1.08, "PnR_PPP_Def": 1.05, "Star1_Name": "LeBron James", "Star1_PPG": 25.4, "Star1_Line": 24.5, "Star1_Style": "ISO", "Star1_Std": 5.5, "Star2_Name": "Anthony Davis", "Star2_PPG": 24.8, "Star2_Line": 24.5, "Star2_Style": "PNR", "Star2_Std": 5.8},
    {"Team": "Phoenix Suns", "Team_Abbr": "PHX", "Conference": "West", "Division": "Pacífico", "Pace": 98.2, "Off_Rtg": 117.9, "Def_Rtg": 114.4, "3PA_Rate": 0.392, "3P_Pct": 0.382, "TOV_Pct": 14.1, "Iso_PPP_Off": 1.13, "Iso_PPP_Def": 1.02, "PnR_PPP_Off": 1.09, "PnR_PPP_Def": 1.03, "Star1_Name": "Kevin Durant", "Star1_PPG": 27.3, "Star1_Line": 26.5, "Star1_Style": "ISO", "Star1_Std": 5.6, "Star2_Name": "Devin Booker", "Star2_PPG": 26.9, "Star2_Line": 26.5, "Star2_Style": "PNR", "Star2_Std": 6.1},
    {"Team": "Sacramento Kings", "Team_Abbr": "SAC", "Conference": "West", "Division": "Pacífico", "Pace": 100.1, "Off_Rtg": 117.2, "Def_Rtg": 114.9, "3PA_Rate": 0.396, "3P_Pct": 0.368, "TOV_Pct": 12.8, "Iso_PPP_Off": 1.08, "Iso_PPP_Def": 1.04, "PnR_PPP_Off": 1.10, "PnR_PPP_Def": 1.04, "Star1_Name": "De'Aaron Fox", "Star1_PPG": 26.4, "Star1_Line": 25.5, "Star1_Style": "PNR", "Star1_Std": 6.0, "Star2_Name": "DeMar DeRozan", "Star2_PPG": 22.8, "Star2_Line": 21.5, "Star2_Style": "ISO", "Star2_Std": 5.0},
    {"Team": "Atlanta Hawks", "Team_Abbr": "ATL", "Conference": "East", "Division": "Sureste", "Pace": 101.8, "Off_Rtg": 116.4, "Def_Rtg": 117.6, "3PA_Rate": 0.394, "3P_Pct": 0.364, "TOV_Pct": 13.7, "Iso_PPP_Off": 1.03, "Iso_PPP_Def": 1.09, "PnR_PPP_Off": 1.11, "PnR_PPP_Def": 1.10, "Star1_Name": "Trae Young", "Star1_PPG": 25.8, "Star1_Line": 25.5, "Star1_Style": "PNR", "Star1_Std": 6.3, "Star2_Name": "Jalen Johnson", "Star2_PPG": 19.4, "Star2_Line": 18.5, "Star2_Style": "PNR", "Star2_Std": 5.1},
    {"Team": "Charlotte Hornets", "Team_Abbr": "CHA", "Conference": "East", "Division": "Sureste", "Pace": 98.9, "Off_Rtg": 111.2, "Def_Rtg": 117.8, "3PA_Rate": 0.415, "3P_Pct": 0.355, "TOV_Pct": 14.6, "Iso_PPP_Off": 0.98, "Iso_PPP_Def": 1.10, "PnR_PPP_Off": 1.03, "PnR_PPP_Def": 1.10, "Star1_Name": "LaMelo Ball", "Star1_PPG": 26.2, "Star1_Line": 25.5, "Star1_Style": "PNR", "Star1_Std": 6.5, "Star2_Name": "Brandon Miller", "Star2_PPG": 20.8, "Star2_Line": 20.5, "Star2_Style": "ISO", "Star2_Std": 5.5},
    {"Team": "Miami Heat", "Team_Abbr": "MIA", "Conference": "East", "Division": "Sureste", "Pace": 95.8, "Off_Rtg": 113.6, "Def_Rtg": 111.5, "3PA_Rate": 0.388, "3P_Pct": 0.370, "TOV_Pct": 12.4, "Iso_PPP_Off": 1.03, "Iso_PPP_Def": 0.93, "PnR_PPP_Off": 1.01, "PnR_PPP_Def": 0.92, "Star1_Name": "Jimmy Butler", "Star1_PPG": 21.4, "Star1_Line": 20.5, "Star1_Style": "ISO", "Star1_Std": 5.1, "Star2_Name": "Bam Adebayo", "Star2_PPG": 19.6, "Star2_Line": 19.5, "Star2_Style": "PNR", "Star2_Std": 4.6},
    {"Team": "Orlando Magic", "Team_Abbr": "ORL", "Conference": "East", "Division": "Sureste", "Pace": 96.8, "Off_Rtg": 112.9, "Def_Rtg": 109.9, "3PA_Rate": 0.354, "3P_Pct": 0.352, "TOV_Pct": 14.1, "Iso_PPP_Off": 1.02, "Iso_PPP_Def": 0.91, "PnR_PPP_Off": 1.01, "PnR_PPP_Def": 0.91, "Star1_Name": "Paolo Banchero", "Star1_PPG": 25.2, "Star1_Line": 24.5, "Star1_Style": "ISO", "Star1_Std": 5.7, "Star2_Name": "Franz Wagner", "Star2_PPG": 22.4, "Star2_Line": 21.5, "Star2_Style": "PNR", "Star2_Std": 5.3},
    {"Team": "Washington Wizards", "Team_Abbr": "WAS", "Conference": "East", "Division": "Sureste", "Pace": 102.6, "Off_Rtg": 109.8, "Def_Rtg": 119.4, "3PA_Rate": 0.392, "3P_Pct": 0.347, "TOV_Pct": 15.1, "Iso_PPP_Off": 0.95, "Iso_PPP_Def": 1.14, "PnR_PPP_Off": 0.96, "PnR_PPP_Def": 1.13, "Star1_Name": "Jordan Poole", "Star1_PPG": 20.4, "Star1_Line": 19.5, "Star1_Style": "PNR", "Star1_Std": 6.1, "Star2_Name": "Kyle Kuzma", "Star2_PPG": 18.8, "Star2_Line": 18.5, "Star2_Style": "ISO", "Star2_Std": 5.4},
    {"Team": "Dallas Mavericks", "Team_Abbr": "DAL", "Conference": "West", "Division": "Suroeste", "Pace": 100.2, "Off_Rtg": 118.1, "Def_Rtg": 114.6, "3PA_Rate": 0.441, "3P_Pct": 0.372, "TOV_Pct": 12.3, "Iso_PPP_Off": 1.15, "Iso_PPP_Def": 1.03, "PnR_PPP_Off": 1.12, "PnR_PPP_Def": 1.02, "Star1_Name": "Luka Doncic", "Star1_PPG": 32.4, "Star1_Line": 31.5, "Star1_Style": "ISO", "Star1_Std": 6.6, "Star2_Name": "Kyrie Irving", "Star2_PPG": 25.3, "Star2_Line": 24.5, "Star2_Style": "ISO", "Star2_Std": 5.9},
    {"Team": "Houston Rockets", "Team_Abbr": "HOU", "Conference": "West", "Division": "Suroeste", "Pace": 99.4, "Off_Rtg": 115.2, "Def_Rtg": 111.2, "3PA_Rate": 0.378, "3P_Pct": 0.358, "TOV_Pct": 12.6, "Iso_PPP_Off": 1.04, "Iso_PPP_Def": 0.94, "PnR_PPP_Off": 1.05, "PnR_PPP_Def": 0.93, "Star1_Name": "Alperen Sengun", "Star1_PPG": 21.3, "Star1_Line": 20.5, "Star1_Style": "PNR", "Star1_Std": 5.2, "Star2_Name": "Jalen Green", "Star2_PPG": 21.8, "Star2_Line": 21.5, "Star2_Style": "ISO", "Star2_Std": 6.2},
    {"Team": "Memphis Grizzlies", "Team_Abbr": "MEM", "Conference": "West", "Division": "Suroeste", "Pace": 101.6, "Off_Rtg": 116.1, "Def_Rtg": 113.0, "3PA_Rate": 0.382, "3P_Pct": 0.366, "TOV_Pct": 13.8, "Iso_PPP_Off": 1.06, "Iso_PPP_Def": 0.97, "PnR_PPP_Off": 1.09, "PnR_PPP_Def": 0.96, "Star1_Name": "Ja Morant", "Star1_PPG": 25.6, "Star1_Line": 24.5, "Star1_Style": "PNR", "Star1_Std": 5.9, "Star2_Name": "Jaren Jackson Jr.", "Star2_PPG": 22.1, "Star2_Line": 21.5, "Star2_Style": "ISO", "Star2_Std": 5.4},
    {"Team": "New Orleans Pelicans", "Team_Abbr": "NOP", "Conference": "West", "Division": "Suroeste", "Pace": 98.3, "Off_Rtg": 114.8, "Def_Rtg": 113.9, "3PA_Rate": 0.371, "3P_Pct": 0.369, "TOV_Pct": 13.4, "Iso_PPP_Off": 1.07, "Iso_PPP_Def": 1.01, "PnR_PPP_Off": 1.04, "PnR_PPP_Def": 1.00, "Star1_Name": "Zion Williamson", "Star1_PPG": 23.6, "Star1_Line": 22.5, "Star1_Style": "ISO", "Star1_Std": 5.5, "Star2_Name": "Brandon Ingram", "Star2_PPG": 21.9, "Star2_Line": 21.5, "Star2_Style": "ISO", "Star2_Std": 5.2},
    {"Team": "San Antonio Spurs", "Team_Abbr": "SAS", "Conference": "West", "Division": "Suroeste", "Pace": 100.9, "Off_Rtg": 113.1, "Def_Rtg": 113.4, "3PA_Rate": 0.391, "3P_Pct": 0.359, "TOV_Pct": 14.5, "Iso_PPP_Off": 1.02, "Iso_PPP_Def": 0.96, "PnR_PPP_Off": 1.05, "PnR_PPP_Def": 0.95, "Star1_Name": "Victor Wembanyama", "Star1_PPG": 24.4, "Star1_Line": 23.5, "Star1_Style": "PNR", "Star1_Std": 5.8, "Star2_Name": "Devin Vassell", "Star2_PPG": 18.9, "Star2_Line": 18.5, "Star2_Style": "PNR", "Star2_Std": 5.0},
]

DEFAULT_LINEUPS_30_TEAMS: Dict[str, List[Dict]] = {
    "BOS": [{"Posición": "PG", "Jugador": "Jrue Holiday", "Minutos_Esperados": 33.0, "O_EPM": 1.4, "D_EPM": 2.3}, {"Posición": "SG", "Jugador": "Derrick White", "Minutos_Esperados": 34.0, "O_EPM": 2.6, "D_EPM": 2.4}, {"Posición": "SF", "Jugador": "Jaylen Brown", "Minutos_Esperados": 35.0, "O_EPM": 2.8, "D_EPM": 0.9}, {"Posición": "PF", "Jugador": "Jayson Tatum", "Minutos_Esperados": 36.0, "O_EPM": 4.8, "D_EPM": 1.5}, {"Posición": "C", "Jugador": "Kristaps Porzingis", "Minutos_Esperados": 30.0, "O_EPM": 2.9, "D_EPM": 2.1}, {"Posición": "BENCH", "Jugador": "Pritchard / Horford / Hauser", "Minutos_Esperados": 72.0, "O_EPM": 0.9, "D_EPM": 0.6}],
    "BKN": [{"Posición": "PG", "Jugador": "Dennis Schroder", "Minutos_Esperados": 32.0, "O_EPM": 0.5, "D_EPM": -0.6}, {"Posición": "SG", "Jugador": "Cam Thomas", "Minutos_Esperados": 34.0, "O_EPM": 1.6, "D_EPM": -1.8}, {"Posición": "SF", "Jugador": "Cameron Johnson", "Minutos_Esperados": 31.0, "O_EPM": 1.1, "D_EPM": -0.2}, {"Posición": "PF", "Jugador": "Dorian Finney-Smith", "Minutos_Esperados": 29.0, "O_EPM": -1.2, "D_EPM": 0.8}, {"Posición": "C", "Jugador": "Nic Claxton", "Minutos_Esperados": 30.0, "O_EPM": -0.5, "D_EPM": 1.9}, {"Posición": "BENCH", "Jugador": "Simmons / Clowney / Wilson", "Minutos_Esperados": 84.0, "O_EPM": -1.6, "D_EPM": -0.6}],
    "NYK": [{"Posición": "PG", "Jugador": "Jalen Brunson", "Minutos_Esperados": 36.0, "O_EPM": 5.4, "D_EPM": -0.4}, {"Posición": "SG", "Jugador": "Mikal Bridges", "Minutos_Esperados": 35.0, "O_EPM": 1.2, "D_EPM": 0.9}, {"Posición": "SF", "Jugador": "OG Anunoby", "Minutos_Esperados": 35.0, "O_EPM": 1.0, "D_EPM": 3.1}, {"Posición": "PF", "Jugador": "Josh Hart", "Minutos_Esperados": 34.0, "O_EPM": 0.4, "D_EPM": 1.5}, {"Posición": "C", "Jugador": "Karl-Anthony Towns", "Minutos_Esperados": 34.0, "O_EPM": 3.1, "D_EPM": 0.5}, {"Posición": "BENCH", "Jugador": "McBride / Achiuwa / Payne", "Minutos_Esperados": 66.0, "O_EPM": -0.4, "D_EPM": 0.6}],
    "PHI": [{"Posición": "PG", "Jugador": "Tyrese Maxey", "Minutos_Esperados": 36.0, "O_EPM": 3.8, "D_EPM": -0.3}, {"Posición": "SG", "Jugador": "Kelly Oubre Jr.", "Minutos_Esperados": 31.0, "O_EPM": -0.3, "D_EPM": 0.4}, {"Posición": "SF", "Jugador": "Paul George", "Minutos_Esperados": 34.0, "O_EPM": 2.9, "D_EPM": 1.6}, {"Posición": "PF", "Jugador": "Caleb Martin", "Minutos_Esperados": 29.0, "O_EPM": -0.6, "D_EPM": 0.7}, {"Posición": "C", "Jugador": "Joel Embiid", "Minutos_Esperados": 34.0, "O_EPM": 5.6, "D_EPM": 3.2}, {"Posición": "BENCH", "Jugador": "Lowry / Drummond / McCain", "Minutos_Esperados": 76.0, "O_EPM": -1.4, "D_EPM": -0.2}],
    "TOR": [{"Posición": "PG", "Jugador": "Immanuel Quickley", "Minutos_Esperados": 33.0, "O_EPM": 1.9, "D_EPM": -0.3}, {"Posición": "SG", "Jugador": "Gradey Dick", "Minutos_Esperados": 29.0, "O_EPM": 0.2, "D_EPM": -1.4}, {"Posición": "SF", "Jugador": "RJ Barrett", "Minutos_Esperados": 33.0, "O_EPM": 1.1, "D_EPM": -0.9}, {"Posición": "PF", "Jugador": "Scottie Barnes", "Minutos_Esperados": 35.0, "O_EPM": 1.8, "D_EPM": 1.2}, {"Posición": "C", "Jugador": "Jakob Poeltl", "Minutos_Esperados": 30.0, "O_EPM": 0.3, "D_EPM": 1.4}, {"Posición": "BENCH", "Jugador": "Boucher / Agbaji / Mitchell", "Minutos_Esperados": 80.0, "O_EPM": -1.8, "D_EPM": -1.1}],
    "CHI": [{"Posición": "PG", "Jugador": "Josh Giddey", "Minutos_Esperados": 31.0, "O_EPM": 0.6, "D_EPM": -0.8}, {"Posición": "SG", "Jugador": "Coby White", "Minutos_Esperados": 34.0, "O_EPM": 1.7, "D_EPM": -0.7}, {"Posición": "SF", "Jugador": "Zach LaVine", "Minutos_Esperados": 34.0, "O_EPM": 2.1, "D_EPM": -1.1}, {"Posición": "PF", "Jugador": "Patrick Williams", "Minutos_Esperados": 28.0, "O_EPM": -1.3, "D_EPM": 0.5}, {"Posición": "C", "Jugador": "Nikola Vucevic", "Minutos_Esperados": 32.0, "O_EPM": 0.8, "D_EPM": -0.6}, {"Posición": "BENCH", "Jugador": "Dosunmu / Ball / Smith", "Minutos_Esperados": 81.0, "O_EPM": -1.1, "D_EPM": -0.2}],
    "CLE": [{"Posición": "PG", "Jugador": "Darius Garland", "Minutos_Esperados": 32.0, "O_EPM": 2.4, "D_EPM": -0.2}, {"Posición": "SG", "Jugador": "Donovan Mitchell", "Minutos_Esperados": 35.0, "O_EPM": 4.9, "D_EPM": 1.1}, {"Posición": "SF", "Jugador": "Max Strus", "Minutos_Esperados": 29.0, "O_EPM": 0.8, "D_EPM": 0.4}, {"Posición": "PF", "Jugador": "Evan Mobley", "Minutos_Esperados": 33.0, "O_EPM": 1.6, "D_EPM": 3.0}, {"Posición": "C", "Jugador": "Jarrett Allen", "Minutos_Esperados": 32.0, "O_EPM": 1.5, "D_EPM": 2.4}, {"Posición": "BENCH", "Jugador": "LeVert / Jerome / Wade", "Minutos_Esperados": 79.0, "O_EPM": 0.7, "D_EPM": 0.6}],
    "DET": [{"Posición": "PG", "Jugador": "Cade Cunningham", "Minutos_Esperados": 35.0, "O_EPM": 2.2, "D_EPM": -0.9}, {"Posición": "SG", "Jugador": "Jaden Ivey", "Minutos_Esperados": 31.0, "O_EPM": 0.3, "D_EPM": -1.4}, {"Posición": "SF", "Jugador": "Tim Hardaway Jr.", "Minutos_Esperados": 28.0, "O_EPM": -0.4, "D_EPM": -1.5}, {"Posición": "PF", "Jugador": "Tobias Harris", "Minutos_Esperados": 32.0, "O_EPM": 0.4, "D_EPM": -0.3}, {"Posición": "C", "Jugador": "Jalen Duren", "Minutos_Esperados": 30.0, "O_EPM": 0.2, "D_EPM": -0.6}, {"Posición": "BENCH", "Jugador": "Beasley / Thompson / Stewart", "Minutos_Esperados": 84.0, "O_EPM": -2.2, "D_EPM": -0.7}],
    "IND": [{"Posición": "PG", "Jugador": "Tyrese Haliburton", "Minutos_Esperados": 35.0, "O_EPM": 5.1, "D_EPM": -1.3}, {"Posición": "SG", "Jugador": "Andrew Nembhard", "Minutos_Esperados": 30.0, "O_EPM": 0.8, "D_EPM": 0.9}, {"Posición": "SF", "Jugador": "Aaron Nesmith", "Minutos_Esperados": 29.0, "O_EPM": 0.7, "D_EPM": 0.6}, {"Posición": "PF", "Jugador": "Pascal Siakam", "Minutos_Esperados": 34.0, "O_EPM": 2.7, "D_EPM": 0.4}, {"Posición": "C", "Jugador": "Myles Turner", "Minutos_Esperados": 30.0, "O_EPM": 1.4, "D_EPM": 1.1}, {"Posición": "BENCH", "Jugador": "Mathurin / McConnell / Toppin", "Minutos_Esperados": 82.0, "O_EPM": 0.8, "D_EPM": -1.2}],
    "MIL": [{"Posición": "PG", "Jugador": "Damian Lillard", "Minutos_Esperados": 35.0, "O_EPM": 3.9, "D_EPM": -1.2}, {"Posición": "SG", "Jugador": "Gary Trent Jr.", "Minutos_Esperados": 30.0, "O_EPM": 0.5, "D_EPM": -0.3}, {"Posición": "SF", "Jugador": "Taurean Prince", "Minutos_Esperados": 28.0, "O_EPM": -0.4, "D_EPM": 0.2}, {"Posición": "PF", "Jugador": "Giannis Antetokounmpo", "Minutos_Esperados": 36.0, "O_EPM": 5.5, "D_EPM": 2.3}, {"Posición": "C", "Jugador": "Brook Lopez", "Minutos_Esperados": 30.0, "O_EPM": 0.3, "D_EPM": 1.8}, {"Posición": "BENCH", "Jugador": "Portis / Wright / Connaughton", "Minutos_Esperados": 81.0, "O_EPM": -0.9, "D_EPM": -0.4}],
    "DEN": [{"Posición": "PG", "Jugador": "Jamal Murray", "Minutos_Esperados": 35.0, "O_EPM": 3.2, "D_EPM": -0.1}, {"Posición": "SG", "Jugador": "Christian Braun", "Minutos_Esperados": 31.0, "O_EPM": 0.3, "D_EPM": 1.1}, {"Posición": "SF", "Jugador": "Michael Porter Jr.", "Minutos_Esperados": 34.0, "O_EPM": 2.0, "D_EPM": 0.3}, {"Posición": "PF", "Jugador": "Aaron Gordon", "Minutos_Esperados": 33.0, "O_EPM": 1.3, "D_EPM": 1.4}, {"Posición": "C", "Jugador": "Nikola Jokic", "Minutos_Esperados": 36.0, "O_EPM": 6.8, "D_EPM": 2.2}, {"Posición": "BENCH", "Jugador": "Westbrook / Watson / Strawther", "Minutos_Esperados": 71.0, "O_EPM": -1.4, "D_EPM": 0.2}],
    "MIN": [{"Posición": "PG", "Jugador": "Mike Conley", "Minutos_Esperados": 29.0, "O_EPM": 1.8, "D_EPM": 0.7}, {"Posición": "SG", "Jugador": "Anthony Edwards", "Minutos_Esperados": 36.0, "O_EPM": 4.2, "D_EPM": 1.4}, {"Posición": "SF", "Jugador": "Jaden McDaniels", "Minutos_Esperados": 31.0, "O_EPM": -0.6, "D_EPM": 2.2}, {"Posición": "PF", "Jugador": "Julius Randle", "Minutos_Esperados": 34.0, "O_EPM": 2.1, "D_EPM": 0.1}, {"Posición": "C", "Jugador": "Rudy Gobert", "Minutos_Esperados": 34.0, "O_EPM": 0.1, "D_EPM": 3.8}, {"Posición": "BENCH", "Jugador": "Reid / DiVincenzo / Alexander-Walker", "Minutos_Esperados": 76.0, "O_EPM": 0.4, "D_EPM": 1.3}],
    "OKC": [{"Posición": "PG", "Jugador": "Shai Gilgeous-Alexander", "Minutos_Esperados": 35.0, "O_EPM": 6.1, "D_EPM": 2.5}, {"Posición": "SG", "Jugador": "Luguentz Dort", "Minutos_Esperados": 30.0, "O_EPM": 0.1, "D_EPM": 2.2}, {"Posición": "SF", "Jugador": "Alex Caruso", "Minutos_Esperados": 28.0, "O_EPM": 0.3, "D_EPM": 3.4}, {"Posición": "PF", "Jugador": "Jalen Williams", "Minutos_Esperados": 34.0, "O_EPM": 3.1, "D_EPM": 1.6}, {"Posición": "C", "Jugador": "Chet Holmgren", "Minutos_Esperados": 32.0, "O_EPM": 2.2, "D_EPM": 3.5}, {"Posición": "BENCH", "Jugador": "Wallace / Joe / Wiggins", "Minutos_Esperados": 81.0, "O_EPM": 0.6, "D_EPM": 1.0}],
    "POR": [{"Posición": "PG", "Jugador": "Scoot Henderson", "Minutos_Esperados": 30.0, "O_EPM": -1.4, "D_EPM": -1.2}, {"Posición": "SG", "Jugador": "Anfernee Simons", "Minutos_Esperados": 34.0, "O_EPM": 1.8, "D_EPM": -2.1}, {"Posición": "SF", "Jugador": "Deni Avdija", "Minutos_Esperados": 31.0, "O_EPM": 0.4, "D_EPM": 0.8}, {"Posición": "PF", "Jugador": "Jerami Grant", "Minutos_Esperados": 33.0, "O_EPM": 0.9, "D_EPM": -0.7}, {"Posición": "C", "Jugador": "Deandre Ayton", "Minutos_Esperados": 31.0, "O_EPM": 0.1, "D_EPM": 0.5}, {"Posición": "BENCH", "Jugador": "Sharpe / Clingan / Camara", "Minutos_Esperados": 81.0, "O_EPM": -2.1, "D_EPM": -0.4}],
    "UTA": [{"Posición": "PG", "Jugador": "Keyonte George", "Minutos_Esperados": 31.0, "O_EPM": 0.5, "D_EPM": -1.9}, {"Posición": "SG", "Jugador": "Collin Sexton", "Minutos_Esperados": 30.0, "O_EPM": 2.1, "D_EPM": -1.5}, {"Posición": "SF", "Jugador": "Lauri Markkanen", "Minutos_Esperados": 34.0, "O_EPM": 3.5, "D_EPM": 0.2}, {"Posición": "PF", "Jugador": "Taylor Hendricks", "Minutos_Esperados": 27.0, "O_EPM": -1.1, "D_EPM": 0.4}, {"Posición": "C", "Jugador": "Walker Kessler", "Minutos_Esperados": 29.0, "O_EPM": -0.2, "D_EPM": 2.3}, {"Posición": "BENCH", "Jugador": "Clarkson / Collins / Filipowski", "Minutos_Esperados": 89.0, "O_EPM": -1.2, "D_EPM": -1.8}],
    "GSW": [{"Posición": "PG", "Jugador": "Stephen Curry", "Minutos_Esperados": 34.0, "O_EPM": 5.6, "D_EPM": -0.3}, {"Posición": "SG", "Jugador": "Brandin Podziemski", "Minutos_Esperados": 29.0, "O_EPM": 0.6, "D_EPM": 0.8}, {"Posición": "SF", "Jugador": "Andrew Wiggins", "Minutos_Esperados": 30.0, "O_EPM": 0.2, "D_EPM": 0.9}, {"Posición": "PF", "Jugador": "Jonathan Kuminga", "Minutos_Esperados": 30.0, "O_EPM": 1.1, "D_EPM": 0.1}, {"Posición": "C", "Jugador": "Draymond Green", "Minutos_Esperados": 30.0, "O_EPM": 0.4, "D_EPM": 3.3}, {"Posición": "BENCH", "Jugador": "Hield / Melton / Jackson-Davis", "Minutos_Esperados": 87.0, "O_EPM": 0.3, "D_EPM": 0.5}],
    "LAC": [{"Posición": "PG", "Jugador": "James Harden", "Minutos_Esperados": 35.0, "O_EPM": 3.7, "D_EPM": 0.1}, {"Posición": "SG", "Jugador": "Norman Powell", "Minutos_Esperados": 33.0, "O_EPM": 2.2, "D_EPM": -0.4}, {"Posición": "SF", "Jugador": "Terance Mann", "Minutos_Esperados": 29.0, "O_EPM": -0.2, "D_EPM": 0.9}, {"Posición": "PF", "Jugador": "Derrick Jones Jr.", "Minutos_Esperados": 28.0, "O_EPM": -0.3, "D_EPM": 1.4}, {"Posición": "C", "Jugador": "Ivica Zubac", "Minutos_Esperados": 32.0, "O_EPM": 1.2, "D_EPM": 1.9}, {"Posición": "BENCH", "Jugador": "Dunn / Coffey / Batum", "Minutos_Esperados": 83.0, "O_EPM": -0.8, "D_EPM": 0.8}],
    "LAL": [{"Posición": "PG", "Jugador": "D'Angelo Russell", "Minutos_Esperados": 31.0, "O_EPM": 1.8, "D_EPM": -1.1}, {"Posición": "SG", "Jugador": "Austin Reaves", "Minutos_Esperados": 34.0, "O_EPM": 1.9, "D_EPM": -0.5}, {"Posición": "SF", "Jugador": "Rui Hachimura", "Minutos_Esperados": 30.0, "O_EPM": 0.6, "D_EPM": -0.6}, {"Posición": "PF", "Jugador": "LeBron James", "Minutos_Esperados": 35.0, "O_EPM": 4.4, "D_EPM": 0.9}, {"Posición": "C", "Jugador": "Anthony Davis", "Minutos_Esperados": 36.0, "O_EPM": 3.1, "D_EPM": 3.7}, {"Posición": "BENCH", "Jugador": "Vincent / Knecht / Hayes", "Minutos_Esperados": 74.0, "O_EPM": -1.6, "D_EPM": -0.4}],
    "PHX": [{"Posición": "PG", "Jugador": "Tyus Jones", "Minutos_Esperados": 30.0, "O_EPM": 1.7, "D_EPM": -0.8}, {"Posición": "SG", "Jugador": "Devin Booker", "Minutos_Esperados": 36.0, "O_EPM": 4.3, "D_EPM": -0.2}, {"Posición": "SF", "Jugador": "Bradley Beal", "Minutos_Esperados": 33.0, "O_EPM": 1.5, "D_EPM": -0.4}, {"Posición": "PF", "Jugador": "Kevin Durant", "Minutos_Esperados": 36.0, "O_EPM": 4.5, "D_EPM": 1.3}, {"Posición": "C", "Jugador": "Jusuf Nurkic", "Minutos_Esperados": 28.0, "O_EPM": 0.2, "D_EPM": 1.5}, {"Posición": "BENCH", "Jugador": "Allen / O'Neale / Plumlee", "Minutos_Esperados": 77.0, "O_EPM": -0.5, "D_EPM": 0.2}],
    "SAC": [{"Posición": "PG", "Jugador": "De'Aaron Fox", "Minutos_Esperados": 36.0, "O_EPM": 3.8, "D_EPM": 0.8}, {"Posición": "SG", "Jugador": "Kevin Huerter", "Minutos_Esperados": 28.0, "O_EPM": 0.6, "D_EPM": -0.7}, {"Posición": "SF", "Jugador": "DeMar DeRozan", "Minutos_Esperados": 35.0, "O_EPM": 2.4, "D_EPM": -0.3}, {"Posición": "PF", "Jugador": "Keegan Murray", "Minutos_Esperados": 34.0, "O_EPM": 0.9, "D_EPM": 1.1}, {"Posición": "C", "Jugador": "Domantas Sabonis", "Minutos_Esperados": 36.0, "O_EPM": 3.3, "D_EPM": 0.9}, {"Posición": "BENCH", "Jugador": "Monk / Ellis / Lyles", "Minutos_Esperados": 71.0, "O_EPM": -0.6, "D_EPM": -0.2}],
    "ATL": [{"Posición": "PG", "Jugador": "Trae Young", "Minutos_Esperados": 36.0, "O_EPM": 4.6, "D_EPM": -2.2}, {"Posición": "SG", "Jugador": "Dyson Daniels", "Minutos_Esperados": 30.0, "O_EPM": -0.5, "D_EPM": 2.4}, {"Posición": "SF", "Jugador": "Zaccharie Risacher", "Minutos_Esperados": 28.0, "O_EPM": -0.4, "D_EPM": 0.2}, {"Posición": "PF", "Jugador": "Jalen Johnson", "Minutos_Esperados": 35.0, "O_EPM": 2.1, "D_EPM": 0.9}, {"Posición": "C", "Jugador": "Clint Capela", "Minutos_Esperados": 27.0, "O_EPM": 0.3, "D_EPM": 1.3}, {"Posición": "BENCH", "Jugador": "Bogdanovic / Hunter / Okongwu", "Minutos_Esperados": 84.0, "O_EPM": 0.2, "D_EPM": -1.3}],
    "CHA": [{"Posición": "PG", "Jugador": "LaMelo Ball", "Minutos_Esperados": 34.0, "O_EPM": 3.6, "D_EPM": -1.4}, {"Posición": "SG", "Jugador": "Josh Green", "Minutos_Esperados": 28.0, "O_EPM": -0.6, "D_EPM": 0.3}, {"Posición": "SF", "Jugador": "Brandon Miller", "Minutos_Esperados": 34.0, "O_EPM": 1.4, "D_EPM": -0.8}, {"Posición": "PF", "Jugador": "Miles Bridges", "Minutos_Esperados": 34.0, "O_EPM": 0.9, "D_EPM": -0.9}, {"Posición": "C", "Jugador": "Mark Williams", "Minutos_Esperados": 29.0, "O_EPM": 0.5, "D_EPM": 1.1}, {"Posición": "BENCH", "Jugador": "Mann / Micic / Richards", "Minutos_Esperados": 81.0, "O_EPM": -2.2, "D_EPM": -1.2}],
    "MIA": [{"Posición": "PG", "Jugador": "Terry Rozier", "Minutos_Esperados": 31.0, "O_EPM": 0.8, "D_EPM": -0.4}, {"Posición": "SG", "Jugador": "Tyler Herro", "Minutos_Esperados": 34.0, "O_EPM": 2.1, "D_EPM": -0.3}, {"Posición": "SF", "Jugador": "Jimmy Butler", "Minutos_Esperados": 34.0, "O_EPM": 3.6, "D_EPM": 1.9}, {"Posición": "PF", "Jugador": "Nikola Jovic", "Minutos_Esperados": 26.0, "O_EPM": 0.2, "D_EPM": 0.4}, {"Posición": "C", "Jugador": "Bam Adebayo", "Minutos_Esperados": 35.0, "O_EPM": 1.1, "D_EPM": 3.2}, {"Posición": "BENCH", "Jugador": "Jaquez Jr. / Robinson / Highsmith", "Minutos_Esperados": 80.0, "O_EPM": -1.1, "D_EPM": 0.9}],
    "ORL": [{"Posición": "PG", "Jugador": "Jalen Suggs", "Minutos_Esperados": 31.0, "O_EPM": 0.9, "D_EPM": 3.1}, {"Posición": "SG", "Jugador": "Kentavious Caldwell-Pope", "Minutos_Esperados": 30.0, "O_EPM": 0.2, "D_EPM": 1.8}, {"Posición": "SF", "Jugador": "Franz Wagner", "Minutos_Esperados": 34.0, "O_EPM": 2.6, "D_EPM": 1.6}, {"Posición": "PF", "Jugador": "Paolo Banchero", "Minutos_Esperados": 35.0, "O_EPM": 2.8, "D_EPM": 0.8}, {"Posición": "C", "Jugador": "Wendell Carter Jr.", "Minutos_Esperados": 29.0, "O_EPM": 0.4, "D_EPM": 1.5}, {"Posición": "BENCH", "Jugador": "Isaac / Anthony / Bitadze", "Minutos_Esperados": 81.0, "O_EPM": -1.6, "D_EPM": 1.4}],
    "WAS": [{"Posición": "PG", "Jugador": "Jordan Poole", "Minutos_Esperados": 32.0, "O_EPM": 0.4, "D_EPM": -2.3}, {"Posición": "SG", "Jugador": "Malcolm Brogdon", "Minutos_Esperados": 28.0, "O_EPM": 0.6, "D_EPM": -0.9}, {"Posición": "SF", "Jugador": "Bilal Coulibaly", "Minutos_Esperados": 30.0, "O_EPM": -1.2, "D_EPM": 0.7}, {"Posición": "PF", "Jugador": "Kyle Kuzma", "Minutos_Esperados": 33.0, "O_EPM": 0.5, "D_EPM": -1.4}, {"Posición": "C", "Jugador": "Alex Sarr", "Minutos_Esperados": 28.0, "O_EPM": -1.6, "D_EPM": 0.6}, {"Posición": "BENCH", "Jugador": "Valanciunas / Kispert / Carrington", "Minutos_Esperados": 89.0, "O_EPM": -1.8, "D_EPM": -1.5}],
    "DAL": [{"Posición": "PG", "Jugador": "Luka Doncic", "Minutos_Esperados": 37.0, "O_EPM": 6.4, "D_EPM": 0.4}, {"Posición": "SG", "Jugador": "Kyrie Irving", "Minutos_Esperados": 35.0, "O_EPM": 4.1, "D_EPM": 0.2}, {"Posición": "SF", "Jugador": "Klay Thompson", "Minutos_Esperados": 30.0, "O_EPM": 1.1, "D_EPM": -0.5}, {"Posición": "PF", "Jugador": "P.J. Washington", "Minutos_Esperados": 32.0, "O_EPM": 0.2, "D_EPM": 1.2}, {"Posición": "C", "Jugador": "Daniel Gafford", "Minutos_Esperados": 26.0, "O_EPM": 1.0, "D_EPM": 1.6}, {"Posición": "BENCH", "Jugador": "Lively II / Grimes / Marshall", "Minutos_Esperados": 80.0, "O_EPM": -0.6, "D_EPM": 0.7}],
    "HOU": [{"Posición": "PG", "Jugador": "Fred VanVleet", "Minutos_Esperados": 35.0, "O_EPM": 2.3, "D_EPM": 1.6}, {"Posición": "SG", "Jugador": "Jalen Green", "Minutos_Esperados": 33.0, "O_EPM": 1.5, "D_EPM": -0.2}, {"Posición": "SF", "Jugador": "Dillon Brooks", "Minutos_Esperados": 31.0, "O_EPM": -0.6, "D_EPM": 1.4}, {"Posición": "PF", "Jugador": "Jabari Smith Jr.", "Minutos_Esperados": 32.0, "O_EPM": 0.3, "D_EPM": 1.3}, {"Posición": "C", "Jugador": "Alperen Sengun", "Minutos_Esperados": 33.0, "O_EPM": 3.1, "D_EPM": 1.1}, {"Posición": "BENCH", "Jugador": "Thompson / Eason / Sheppard", "Minutos_Esperados": 76.0, "O_EPM": -0.8, "D_EPM": 1.5}],
    "MEM": [{"Posición": "PG", "Jugador": "Ja Morant", "Minutos_Esperados": 34.0, "O_EPM": 3.9, "D_EPM": -0.3}, {"Posición": "SG", "Jugador": "Desmond Bane", "Minutos_Esperados": 34.0, "O_EPM": 2.8, "D_EPM": 0.6}, {"Posición": "SF", "Jugador": "Marcus Smart", "Minutos_Esperados": 29.0, "O_EPM": -0.4, "D_EPM": 1.9}, {"Posición": "PF", "Jugador": "Jaren Jackson Jr.", "Minutos_Esperados": 32.0, "O_EPM": 1.5, "D_EPM": 2.8}, {"Posición": "C", "Jugador": "Zach Edey", "Minutos_Esperados": 26.0, "O_EPM": 0.4, "D_EPM": 1.1}, {"Posición": "BENCH", "Jugador": "Aldama / Pippen Jr. / Clarke", "Minutos_Esperados": 85.0, "O_EPM": -0.6, "D_EPM": 0.3}],
    "NOP": [{"Posición": "PG", "Jugador": "Dejounte Murray", "Minutos_Esperados": 34.0, "O_EPM": 1.7, "D_EPM": 0.9}, {"Posición": "SG", "Jugador": "CJ McCollum", "Minutos_Esperados": 32.0, "O_EPM": 1.9, "D_EPM": -0.6}, {"Posición": "SF", "Jugador": "Brandon Ingram", "Minutos_Esperados": 34.0, "O_EPM": 2.2, "D_EPM": -0.2}, {"Posición": "PF", "Jugador": "Herbert Jones", "Minutos_Esperados": 31.0, "O_EPM": 0.2, "D_EPM": 2.9}, {"Posición": "C", "Jugador": "Zion Williamson", "Minutos_Esperados": 33.0, "O_EPM": 3.4, "D_EPM": 0.4}, {"Posición": "BENCH", "Jugador": "Murphy III / Alvarado / Missi", "Minutos_Esperados": 76.0, "O_EPM": -0.9, "D_EPM": 0.2}],
    "SAS": [{"Posición": "PG", "Jugador": "Chris Paul", "Minutos_Esperados": 28.0, "O_EPM": 1.6, "D_EPM": 0.4}, {"Posición": "SG", "Jugador": "Devin Vassell", "Minutos_Esperados": 33.0, "O_EPM": 1.9, "D_EPM": 0.1}, {"Posición": "SF", "Jugador": "Harrison Barnes", "Minutos_Esperados": 29.0, "O_EPM": 0.5, "D_EPM": -0.4}, {"Posición": "PF", "Jugador": "Jeremy Sochan", "Minutos_Esperados": 30.0, "O_EPM": -0.8, "D_EPM": 0.7}, {"Posición": "C", "Jugador": "Victor Wembanyama", "Minutos_Esperados": 34.0, "O_EPM": 2.9, "D_EPM": 4.5}, {"Posición": "BENCH", "Jugador": "Johnson / Castle / Tre Jones", "Minutos_Esperados": 86.0, "O_EPM": -1.4, "D_EPM": -0.1}],
}

# Alias de respaldo base (solo usado al inicializar sesión por primera vez o hacer reset explícito)
QUINTETOS_BASE = DEFAULT_LINEUPS_30_TEAMS
DEFAULT_TEAM_META_BY_ABBR: Dict[str, Dict] = {d["Team_Abbr"]: d for d in DEFAULT_CSV_DATA}


def _normal_cdf(x: float) -> float:
    """Función de distribución acumulada normal estándar Φ(x) exacta vía math.erf."""
    return 0.5 * (1.0 + math.erf(x / math.sqrt(2.0)))


def _poisson_cdf(k_max: int, lam: float) -> float:
    """
    Calcula P(X <= k_max) para una distribución de Poisson con media esperada λ > 0
    utilizando log-factoriales (math.lgamma) para evitar desbordamiento numérico.
    """
    if lam <= 0.0:
        return 1.0 if k_max >= 0 else 0.0
    if k_max < 0:
        return 0.0
    total_prob = 0.0
    log_lam = math.log(lam)
    for k in range(int(k_max) + 1):
        log_pmf = -lam + k * log_lam - math.lgamma(k + 1)
        total_prob += math.exp(log_pmf)
    return float(np.clip(total_prob, 0.001, 0.999))


def prob_to_fair_odds(prob: float) -> Tuple[str, str]:
    """
    Convierte una probabilidad (0-1) a cuota justa estrictamente en FORMATO DECIMAL
    (ej. '1.91', '2.15'). Retorna tupla (dec_str, dec_str) por compatibilidad modular.
    """
    p = float(np.clip(prob, 0.01, 0.99))
    dec_odds = 1.0 / p
    dec_str = f"{dec_odds:.2f}"
    return dec_str, dec_str


def round_to_half_point(val: float) -> float:
    """Redondea cualquier línea proyectada al medio punto (.5) más cercano estilo Sportsbook."""
    return float(np.floor(val) + 0.5)


def _derivar_metricas_base_jugador(
    player_name: str,
    pos: str,
    minutos_base: float,
    o_epm: float,
    d_epm: float,
    team_abbr: str = "",
) -> Tuple[float, float, float]:
    """
    Deriva PPG_Base, RPG_Base y APG_Base realistas para cualquier jugador en función de su
    nombre (si es estrella registrada), posición (PG, SG, SF, PF, C) y ratings O_EPM / D_EPM.
    """
    p_low = str(player_name).strip().lower()
    pos_up = str(pos).strip().upper()
    t_meta = DEFAULT_TEAM_META_BY_ABBR.get(str(team_abbr).strip().upper(), {})

    if pos_up in {"BENCH", "BANCA"}:
        return 0.0, 0.0, 0.0

    s1_name = str(t_meta.get("Star1_Name", "")).strip().lower()
    s2_name = str(t_meta.get("Star2_Name", "")).strip().lower()
    if p_low == s1_name and "Star1_PPG" in t_meta:
        ppg_base = float(t_meta["Star1_PPG"])
    elif p_low == s2_name and "Star2_PPG" in t_meta:
        ppg_base = float(t_meta["Star2_PPG"])
    else:
        ppg_base = max(8.5, (minutos_base * 0.52) + (o_epm * 1.55))

    if p_low in STAR_BASE_REB_AST_MAP:
        rpg_base, apg_base = STAR_BASE_REB_AST_MAP[p_low]
    else:
        min_scale = max(minutos_base, 18.0) / 34.0
        if pos_up == "C":
            rpg_base = (9.6 + max(0.0, d_epm) * 0.65) * min_scale
            apg_base = (2.6 + max(0.0, o_epm) * 0.45) * min_scale
        elif pos_up == "PF":
            rpg_base = (7.4 + max(0.0, d_epm) * 0.55) * min_scale
            apg_base = (3.2 + max(0.0, o_epm) * 0.45) * min_scale
        elif pos_up == "SF":
            rpg_base = (5.6 + max(0.0, d_epm) * 0.40) * min_scale
            apg_base = (3.6 + max(0.0, o_epm) * 0.50) * min_scale
        elif pos_up == "SG":
            rpg_base = (4.3 + max(0.0, d_epm) * 0.30) * min_scale
            apg_base = (4.4 + max(0.0, o_epm) * 0.55) * min_scale
        else:
            rpg_base = (4.0 + max(0.0, d_epm) * 0.25) * min_scale
            apg_base = (6.8 + max(0.0, o_epm) * 0.70) * min_scale

    return round(float(ppg_base), 1), round(float(rpg_base), 1), round(float(apg_base), 1)


def build_default_lineups_dict() -> Dict[str, pd.DataFrame]:
    """
    Construye el diccionario inicial de DataFrames con los quintetos titulares de las 30 franquicias,
    incluyendo las métricas base requeridas ('PPG_Base', 'RPG_Base', 'APG_Base', 'Minutos_Base').
    """
    lineups_dict: Dict[str, pd.DataFrame] = {}
    for abbr, rows in DEFAULT_LINEUPS_30_TEAMS.items():
        enriched_rows: List[Dict[str, Any]] = []
        for r in rows:
            pos_str = str(r["Posición"]).strip().upper()
            p_name = str(r["Jugador"]).strip()
            mins_esp = float(r["Minutos_Esperados"])
            o_epm = float(r["O_EPM"])
            d_epm = float(r["D_EPM"])
            ppg_b, rpg_b, apg_b = _derivar_metricas_base_jugador(p_name, pos_str, mins_esp, o_epm, d_epm, abbr)
            row_copy = dict(r)
            row_copy["Minutos_Base"] = round(mins_esp, 1)
            row_copy["Pace_Diff"] = 0.0
            row_copy["PPG_Base"] = ppg_b
            row_copy["RPG_Base"] = rpg_b
            row_copy["APG_Base"] = apg_b
            enriched_rows.append(row_copy)
        lineups_dict[abbr] = pd.DataFrame(enriched_rows)
    return lineups_dict


def obtener_lineup_en_sesion(team_abbr: str, fallback_lineups: Optional[Dict[str, pd.DataFrame]] = None) -> pd.DataFrame:
    """
    Recupera obligatoriamente el quinteto de `st.session_state.lineups_nba` para `team_abbr`.
    Garantiza que las columnas base ('Minutos_Base', 'PPG_Base', 'RPG_Base', 'APG_Base') estén presentes.
    """
    abbr = str(team_abbr).strip().upper()
    source = fallback_lineups
    if source is None and (st is not None and hasattr(st, "session_state")) and "lineups_nba" in st.session_state:
        source = st.session_state.lineups_nba

    df_res: Optional[pd.DataFrame] = None
    if isinstance(source, dict) and abbr in source:
        df_res = source[abbr].copy()
    elif isinstance(source, pd.DataFrame) and not source.empty and "Team_Abbr" in source.columns:
        sub = source[source["Team_Abbr"].astype(str).str.strip().str.upper() == abbr].copy()
        if not sub.empty:
            if "Pos" in sub.columns and "Posición" not in sub.columns:
                sub["Posición"] = sub["Pos"].apply(
                    lambda x: "BENCH" if str(x).strip().upper() in {"BANCA", "BENCH"} else str(x).strip().upper()
                )
            if "Minutos" in sub.columns and "Minutos_Esperados" not in sub.columns:
                sub["Minutos_Esperados"] = pd.to_numeric(sub["Minutos"], errors="coerce").fillna(0.0)
            df_res = sub

    if df_res is None:
        df_res = build_default_lineups_dict()[abbr].copy()

    if not {"Minutos_Base", "PPG_Base", "RPG_Base", "APG_Base"}.issubset(df_res.columns):
        default_lu = {
            str(r["Jugador"]).lower(): r for r in DEFAULT_LINEUPS_30_TEAMS.get(abbr, [])
        }
        mins_base_list, ppg_b_list, rpg_b_list, apg_b_list = [], [], [], []
        for _, r in df_res.iterrows():
            p_name = str(r.get("Jugador", "")).strip()
            pos_str = str(r.get("Posición", r.get("Pos", "PG"))).strip().upper()
            def_r = default_lu.get(p_name.lower(), {})
            m_base = float(r.get("Minutos_Base", def_r.get("Minutos_Esperados", r.get("Minutos_Esperados", 34.0))))
            if m_base <= 0.0:
                m_base = 34.0
            o_epm = float(r.get("O_EPM", 0.0))
            d_epm = float(r.get("D_EPM", 0.0))
            ppg_d, rpg_d, apg_d = _derivar_metricas_base_jugador(p_name, pos_str, m_base, o_epm, d_epm, abbr)
            mins_base_list.append(round(float(r.get("Minutos_Base", m_base)), 1))
            ppg_b_list.append(round(float(r.get("PPG_Base", ppg_d)), 1))
            rpg_b_list.append(round(float(r.get("RPG_Base", rpg_d)), 1))
            apg_b_list.append(round(float(r.get("APG_Base", apg_d)), 1))
        df_res["Minutos_Base"] = mins_base_list
        df_res["PPG_Base"] = ppg_b_list
        df_res["RPG_Base"] = rpg_b_list
        df_res["APG_Base"] = apg_b_list

    return df_res


# ==============================================================================
# SISTEMA DE AUDITORÍA POST-PARTIDO Y CALIBRACIÓN ADAPTATIVA (FILTRO KALMAN)
# ==============================================================================
def auditar_partido_nba(proyeccion: Dict[str, Any], resultado_real: Dict[str, Any]) -> Dict[str, Any]:
    """
    Módulo de Auditoría Post-Partido para NBA:
    Calcula los residuales entre la proyección previa del modelo Monte Carlo y el resultado real:
      - `Error_Away`: Puntos reales anotados por la visita menos puntos proyectados (`score_away_real - exp_away_pts`)
      - `Error_Home`: Puntos reales anotados por el local menos puntos proyectados (`score_home_real - exp_home_pts`)
      - `Error_Spread`: Diferencial real del local (`margen_real_home`) menos diferencial proyectado (`exp_margin_home`)
      - `Error_Total`: Total de puntos reales (`total_real`) menos total proyectado (`exp_total_pts`)
      - `Error_Pace`: Ritmo real estimado (`pace_real`) menos posesiones proyectadas (`exp_possessions`)
    """
    exp_away = float(proyeccion.get("exp_away_pts", proyeccion.get("pts_away_modelo", 112.0)))
    exp_home = float(proyeccion.get("exp_home_pts", proyeccion.get("pts_home_modelo", 114.5)))
    exp_total = float(proyeccion.get("exp_total_pts", exp_away + exp_home))
    exp_margin_home = float(proyeccion.get("exp_margin_home", exp_home - exp_away))
    proj_spread_home = float(proyeccion.get("proj_spread_home", -exp_margin_home))
    exp_pace = float(proyeccion.get("exp_possessions", proyeccion.get("pace_modelo", 99.0)))

    real_away = float(resultado_real.get("score_away_real", 0.0))
    real_home = float(resultado_real.get("score_home_real", 0.0))
    real_total = real_away + real_home
    real_margin_home = real_home - real_away
    real_spread_home = -real_margin_home
    real_pace = float(
        resultado_real.get("pace_real")
        if resultado_real.get("pace_real") is not None and float(resultado_real.get("pace_real", 0.0)) > 0
        else (real_total / 2.25 if real_total > 0 else exp_pace)
    )

    error_away = round(real_away - exp_away, 2)
    error_home = round(real_home - exp_home, 2)
    error_spread = round(real_margin_home - exp_margin_home, 2)
    error_total = round(real_total - exp_total, 2)
    error_pace = round(real_pace - exp_pace, 2)

    ganador_proy = str(proyeccion.get("home_abbr", "HOME")) if exp_home >= exp_away else str(proyeccion.get("away_abbr", "AWAY"))
    ganador_real = str(proyeccion.get("home_abbr", "HOME")) if real_home >= real_away else str(proyeccion.get("away_abbr", "AWAY"))
    acierto_ganador = bool(ganador_proy == ganador_real)

    return {
        "away_abbr": str(proyeccion.get("away_abbr", resultado_real.get("away_abbr", "AWAY"))),
        "home_abbr": str(proyeccion.get("home_abbr", resultado_real.get("home_abbr", "HOME"))),
        "exp_away_pts": round(exp_away, 1),
        "exp_home_pts": round(exp_home, 1),
        "exp_total_pts": round(exp_total, 1),
        "exp_margin_home": round(exp_margin_home, 1),
        "proj_spread_home": round(proj_spread_home, 1),
        "exp_pace": round(exp_pace, 1),
        "score_away_real": round(real_away, 1),
        "score_home_real": round(real_home, 1),
        "total_real": round(real_total, 1),
        "margen_real_home": round(real_margin_home, 1),
        "spread_real_home": round(real_spread_home, 1),
        "pace_real": round(real_pace, 1),
        "Error_Away": error_away,
        "Error_Home": error_home,
        "Error_Spread": error_spread,
        "Error_Total": error_total,
        "Error_Pace": error_pace,
        "mae_pts": round((abs(error_away) + abs(error_home)) / 2.0, 2),
        "acierto_ganador": acierto_ganador,
    }


def calibrar_ratings_nba(
    df_equipos: pd.DataFrame,
    away_abbr: str,
    home_abbr: str,
    auditoria: Dict[str, Any],
    learning_rate: float = 0.08,
    lineups_dict: Optional[Dict[str, pd.DataFrame]] = None,
) -> pd.DataFrame:
    """
    Calibración Adaptativa tipo Filtro Kalman / Feedback Loop para NBA:
      1. Ajusta el 'Pace' (y 'Base_Pace') de ambos equipos con el residual de ritmo (`Error_Pace`).
      2. Ajusta 'Off_Rtg' y 'Def_Rtg' (y sus columnas base) en el DataFrame de equipos según
         los puntos reales anotados/permitidos frente a la proyección (`Error_Away`, `Error_Home`).
      3. Aplica límites de seguridad (clipping) para mantener los ratings en rangos lógicos de la NBA moderna:
         - Pace: [92.0, 106.5]
         - Off_Rtg: [104.0, 126.5]
         - Def_Rtg: [104.0, 124.5]
      4. Si se proporciona `lineups_dict` (o existe `st.session_state.lineups_nba`), sincroniza también
         el delta de `O_EPM`, `D_EPM` y `Pace_Diff` de los titulares para que el motor Bottom-Up y el
         Excel calibrado reflejen exactamente los nuevos ratings.
    """
    df_cal = df_equipos.copy()
    a_code = str(away_abbr).strip().upper()
    h_code = str(home_abbr).strip().upper()
    lr = float(np.clip(learning_rate, 0.01, 0.35))

    err_away = float(auditoria.get("Error_Away", 0.0))
    err_home = float(auditoria.get("Error_Home", 0.0))
    err_pace = float(auditoria.get("Error_Pace", 0.0))
    pace_ref = max(float(auditoria.get("pace_real", 99.0)), 88.0)

    # Convertir los errores de puntos del partido a escala por 100 posesiones
    err_away_per100 = err_away * (100.0 / pace_ref)
    err_home_per100 = err_home * (100.0 / pace_ref)

    # Deltas Kalman:
    # - Si Away anotó más de lo esperado (err_away > 0): sube Off_Rtg de Away y sube Def_Rtg de Home (peor defensa)
    # - Si Home anotó más de lo esperado (err_home > 0): sube Off_Rtg de Home y sube Def_Rtg de Away (peor defensa)
    delta_away_off = lr * err_away_per100
    delta_home_def = lr * err_away_per100

    delta_home_off = lr * err_home_per100
    delta_away_def = lr * err_home_per100

    delta_pace = lr * err_pace

    away_mask = df_cal["Team_Abbr"].astype(str).str.strip().str.upper() == a_code
    home_mask = df_cal["Team_Abbr"].astype(str).str.strip().str.upper() == h_code

    if away_mask.any():
        idx_a = df_cal[away_mask].index[0]
        new_a_pace = round(float(np.clip(float(df_cal.loc[idx_a, "Pace"]) + delta_pace, *PACE_CLIP_RANGE)), 1)
        new_a_off = round(float(np.clip(float(df_cal.loc[idx_a, "Off_Rtg"]) + delta_away_off, *ORTG_CLIP_RANGE)), 1)
        new_a_def = round(float(np.clip(float(df_cal.loc[idx_a, "Def_Rtg"]) + delta_away_def, *DRTG_CLIP_RANGE)), 1)

        df_cal.loc[idx_a, "Pace"] = new_a_pace
        df_cal.loc[idx_a, "Off_Rtg"] = new_a_off
        df_cal.loc[idx_a, "Def_Rtg"] = new_a_def
        if "Base_Pace" in df_cal.columns:
            df_cal.loc[idx_a, "Base_Pace"] = new_a_pace
        if "Base_Off_Rtg" in df_cal.columns:
            df_cal.loc[idx_a, "Base_Off_Rtg"] = new_a_off
        if "Base_Def_Rtg" in df_cal.columns:
            df_cal.loc[idx_a, "Base_Def_Rtg"] = new_a_def

    if home_mask.any():
        idx_h = df_cal[home_mask].index[0]
        new_h_pace = round(float(np.clip(float(df_cal.loc[idx_h, "Pace"]) + delta_pace, *PACE_CLIP_RANGE)), 1)
        new_h_off = round(float(np.clip(float(df_cal.loc[idx_h, "Off_Rtg"]) + delta_home_off, *ORTG_CLIP_RANGE)), 1)
        new_h_def = round(float(np.clip(float(df_cal.loc[idx_h, "Def_Rtg"]) + delta_home_def, *DRTG_CLIP_RANGE)), 1)

        df_cal.loc[idx_h, "Pace"] = new_h_pace
        df_cal.loc[idx_h, "Off_Rtg"] = new_h_off
        df_cal.loc[idx_h, "Def_Rtg"] = new_h_def
        if "Base_Pace" in df_cal.columns:
            df_cal.loc[idx_h, "Base_Pace"] = new_h_pace
        if "Base_Off_Rtg" in df_cal.columns:
            df_cal.loc[idx_h, "Base_Off_Rtg"] = new_h_off
        if "Base_Def_Rtg" in df_cal.columns:
            df_cal.loc[idx_h, "Base_Def_Rtg"] = new_h_def

    # Sincronizar ajustes Kalman hacia los quintetos en sesión para que el motor Bottom-Up y el Excel exportado los preserven
    target_lineups = lineups_dict
    if target_lineups is None and (st is not None and hasattr(st, "session_state")) and "lineups_nba" in st.session_state:
        target_lineups = st.session_state.lineups_nba

    if isinstance(target_lineups, dict):
        synergy_scale = 0.62
        for t_code, d_off, d_def, d_pace in [
            (a_code, delta_away_off, delta_away_def, delta_pace),
            (h_code, delta_home_off, delta_home_def, delta_pace),
        ]:
            if t_code in target_lineups:
                lu_df = target_lineups[t_code].copy()
                # En calcular_ratings_lineup: proj_off_rtg = 115.5 + sum(slots * O_EPM) * 0.62
                # y proj_def_rtg = 115.5 - sum(slots * D_EPM) * 0.62 (con sum(slots) = 5.0)
                o_epm_step = d_off / (5.0 * synergy_scale)
                d_epm_step = -d_def / (5.0 * synergy_scale)
                pace_step = d_pace / 5.0

                lu_df["O_EPM"] = (pd.to_numeric(lu_df["O_EPM"], errors="coerce").fillna(0.0) + o_epm_step).clip(-8.0, 10.0).round(2)
                lu_df["D_EPM"] = (pd.to_numeric(lu_df["D_EPM"], errors="coerce").fillna(0.0) + d_epm_step).clip(-8.0, 10.0).round(2)
                if "Pace_Diff" not in lu_df.columns:
                    lu_df["Pace_Diff"] = 0.0
                lu_df["Pace_Diff"] = (pd.to_numeric(lu_df["Pace_Diff"], errors="coerce").fillna(0.0) + pace_step).clip(-4.0, 4.0).round(2)
                target_lineups[t_code] = lu_df

    return df_cal


# ==============================================================================
# MOTOR DE PROYECCIÓN DE PLAYER PROPS (PTS, REB, AST) Y VALOR +EV (POISSON DECIMAL)
# ==============================================================================
def evaluar_mercado_prop_poisson(
    mercado: str,
    lambda_esperada: float,
    linea_mercado: float,
    cuota_over_decimal: float = 1.91,
    cuota_under_decimal: Optional[float] = None,
    umbral_ev_pct: float = 5.0,
) -> Dict[str, Any]:
    """
    Evalúa un mercado individual de Player Prop (PTS, REB o AST) frente a una línea de sportsbook
    usando distribución de Poisson con media esperada λ y cuotas en formato decimal.
    """
    lam = max(0.2, float(lambda_esperada))
    linea = round(float(linea_mercado), 1)
    c_over = round(max(1.01, float(cuota_over_decimal)), 2)
    c_under = round(max(1.01, float(cuota_under_decimal if cuota_under_decimal is not None else cuota_over_decimal)), 2)

    k_floor = int(math.floor(linea))
    if abs(linea - round(linea)) < 1e-6 and k_floor >= 1:
        p_under_raw = _poisson_cdf(k_floor - 1, lam)
        p_over_raw = 1.0 - _poisson_cdf(k_floor, lam)
        denom = max(p_over_raw + p_under_raw, 0.01)
        prob_over = float(np.clip(p_over_raw / denom, 0.01, 0.99))
        prob_under = float(np.clip(1.0 - prob_over, 0.01, 0.99))
    else:
        prob_under = float(np.clip(_poisson_cdf(k_floor, lam), 0.01, 0.99))
        prob_over = float(np.clip(1.0 - prob_under, 0.01, 0.99))

    fair_odd_over = round(1.0 / prob_over, 2)
    fair_odd_under = round(1.0 / prob_under, 2)

    ev_over_pct = round(((prob_over * c_over) - 1.0) * 100.0, 1)
    ev_under_pct = round(((prob_under * c_under) - 1.0) * 100.0, 1)

    if ev_over_pct >= float(umbral_ev_pct):
        pick_side = "OVER"
        ev_display = ev_over_pct
        cuota_activa = c_over
        fair_odd_activa = fair_odd_over
        prob_activa = round(prob_over * 100.0, 1)
        badge_text = f"[⚡ VALOR: OVER {linea:.1f} @ {c_over:.2f} ({ev_display:+.1f}% EV)]"
        badge_color = "#00ff88"
        tiene_valor = True
    elif ev_under_pct >= float(umbral_ev_pct) or (ev_over_pct <= -float(umbral_ev_pct) and lam < linea):
        pick_side = "UNDER"
        ev_display = ev_under_pct if ev_under_pct >= float(umbral_ev_pct) else round(abs(ev_over_pct), 1)
        cuota_activa = c_under
        fair_odd_activa = fair_odd_under
        prob_activa = round(prob_under * 100.0, 1)
        badge_text = f"[⚡ VALOR: UNDER {linea:.1f} @ {c_under:.2f} (+{abs(ev_display):.1f}% EV)]"
        badge_color = "#ff3366"
        tiene_valor = True
    else:
        pick_side = "LÍNEA JUSTA"
        ev_display = ev_over_pct if prob_over >= prob_under else ev_under_pct
        cuota_activa = c_over
        fair_odd_activa = fair_odd_over if prob_over >= prob_under else fair_odd_under
        prob_activa = round(max(prob_over, prob_under) * 100.0, 1)
        badge_text = "[LÍNEA JUSTA]"
        badge_color = "#555555"
        tiene_valor = False

    return {
        "mercado": str(mercado).upper(),
        "proyeccion": round(lam, 1),
        "lambda": round(lam, 2),
        "linea": linea,
        "cuota": cuota_activa,
        "cuota_over": c_over,
        "cuota_under": c_under,
        "prob_over": round(prob_over * 100.0, 1),
        "prob_under": round(prob_under * 100.0, 1),
        "fair_odd_over": fair_odd_over,
        "fair_odd_under": fair_odd_under,
        "fair_odd": fair_odd_activa,
        "ev_over_pct": ev_over_pct,
        "ev_under_pct": ev_under_pct,
        "ev_pct": ev_display,
        "pick": pick_side,
        "prob_pick": prob_activa,
        "badge": badge_text,
        "badge_color": badge_color,
        "valor": tiene_valor,
    }


def proyectar_props_jugador(
    jugador_dict: Dict[str, Any],
    minutos_proyectados: float,
    pace_partido: float,
    pace_promedio_liga: float = 99.5,
    def_rating_rival: float = 115.0,
    lineas_mercado: Optional[Dict[str, float]] = None,
    cuotas_mercado: Optional[Dict[str, float]] = None,
) -> Dict[str, Any]:
    """
    Motor Cuantitativo de Proyección de Player Props (PTS, REB, AST) y Auditoría de Valor +EV en Decimal.
    """
    nombre = str(jugador_dict.get("Jugador", jugador_dict.get("player_name", "Jugador"))).strip()
    team_abbr = str(jugador_dict.get("Team_Abbr", jugador_dict.get("team_abbr", "NBA"))).strip().upper()
    pos_str = str(jugador_dict.get("Posición", jugador_dict.get("Pos", "SF"))).strip().upper()
    o_epm = float(jugador_dict.get("O_EPM", 2.0))
    d_epm = float(jugador_dict.get("D_EPM", 0.5))

    minutos_base = float(jugador_dict.get("Minutos_Base", 34.0))
    if minutos_base <= 0.5:
        minutos_base = 34.0

    ppg_def, rpg_def, apg_def = _derivar_metricas_base_jugador(
        nombre, pos_str, minutos_base, o_epm, d_epm, team_abbr
    )
    ppg_base = float(jugador_dict.get("PPG_Base", jugador_dict.get("base_ppg", ppg_def)))
    rpg_base = float(jugador_dict.get("RPG_Base", rpg_def))
    apg_base = float(jugador_dict.get("APG_Base", apg_def))

    mins_proy = max(0.0, float(minutos_proyectados))
    factor_pace = float(pace_partido) / max(float(pace_promedio_liga), 80.0)
    factor_min = mins_proy / minutos_base
    factor_def = float(def_rating_rival) / 115.0

    lambda_pts = max(0.5, ppg_base * factor_min * factor_pace * factor_def)
    lambda_reb = max(0.3, rpg_base * factor_min * factor_pace)
    lambda_ast = max(0.3, apg_base * factor_min * factor_pace * factor_def)

    lineas_in = lineas_mercado or {}
    cuotas_in = cuotas_mercado or {}

    linea_pts = float(
        lineas_in.get(
            "PTS",
            jugador_dict.get("PPG_Linea", jugador_dict.get("prop_line", round_to_half_point(max(0.5, ppg_base - 0.5)))),
        )
    )
    linea_reb = float(
        lineas_in.get(
            "REB",
            jugador_dict.get("RPG_Linea", round_to_half_point(max(0.5, rpg_base - 0.5))),
        )
    )
    linea_ast = float(
        lineas_in.get(
            "AST",
            jugador_dict.get("APG_Linea", round_to_half_point(max(0.5, apg_base - 0.5))),
        )
    )

    cuota_pts = float(cuotas_in.get("PTS", jugador_dict.get("Cuota_PTS", 1.91)))
    cuota_reb = float(cuotas_in.get("REB", jugador_dict.get("Cuota_REB", 1.91)))
    cuota_ast = float(cuotas_in.get("AST", jugador_dict.get("Cuota_AST", 1.91)))

    eval_pts = evaluar_mercado_prop_poisson("PTS", lambda_pts, linea_pts, cuota_over_decimal=cuota_pts)
    eval_reb = evaluar_mercado_prop_poisson("REB", lambda_reb, linea_reb, cuota_over_decimal=cuota_reb)
    eval_ast = evaluar_mercado_prop_poisson("AST", lambda_ast, linea_ast, cuota_over_decimal=cuota_ast)

    return {
        "jugador": nombre,
        "team_abbr": team_abbr,
        "posicion": pos_str,
        "minutos_base": round(minutos_base, 1),
        "minutos_proyectados": round(mins_proy, 1),
        "ppg_base": round(ppg_base, 1),
        "rpg_base": round(rpg_base, 1),
        "apg_base": round(apg_base, 1),
        "factor_pace": round(factor_pace, 3),
        "factor_min": round(factor_min, 3),
        "lambda_pts": round(lambda_pts, 1),
        "lambda_reb": round(lambda_reb, 1),
        "lambda_ast": round(lambda_ast, 1),
        "PTS": eval_pts,
        "REB": eval_reb,
        "AST": eval_ast,
    }


# ==============================================================================
# DETECTORES DE VALOR CUANTITATIVO DE PARTIDO (+EV EN FORMATO DECIMAL)
# ==============================================================================
def evaluar_valor_totales_nba(
    total_modelo: float,
    total_mercado: float,
    cuota_over: float,
    cuota_under: float,
    umbral_pts: float = 3.0,
    prob_over_sim: Optional[float] = None,
    prob_under_sim: Optional[float] = None,
) -> Dict[str, Any]:
    """Auditoría matemática de Totales (Over/Under) en formato estrictamente DECIMAL."""
    t_mod = float(total_modelo)
    t_mkt = round(float(total_mercado), 1)
    c_over = round(float(cuota_over), 2)
    c_under = round(float(cuota_under), 2)
    edge_pts = t_mod - t_mkt

    if prob_over_sim is not None and 0.01 < float(prob_over_sim) < 0.99:
        prob_over = float(prob_over_sim)
    else:
        prob_over = float(np.clip(_normal_cdf(edge_pts / NBA_TOTAL_SIM_STD), 0.05, 0.95))

    if prob_under_sim is not None and 0.01 < float(prob_under_sim) < 0.99:
        prob_under = float(prob_under_sim)
    else:
        prob_under = float(np.clip(1.0 - prob_over, 0.05, 0.95))

    if edge_pts >= float(umbral_pts):
        ev_pct = round(((prob_over * c_over) - 1.0) * 100.0, 1)
        return {
            "pick": "OVER",
            "linea": t_mkt,
            "cuota": c_over,
            "edge_pts": round(edge_pts, 1),
            "ev_pct": ev_pct,
            "prob": round(prob_over * 100.0, 1),
            "badge_color": "#00ff88",
            "valor": True,
        }

    if edge_pts <= -float(umbral_pts):
        ev_pct = round(((prob_under * c_under) - 1.0) * 100.0, 1)
        return {
            "pick": "UNDER",
            "linea": t_mkt,
            "cuota": c_under,
            "edge_pts": round(abs(edge_pts), 1),
            "ev_pct": ev_pct,
            "prob": round(prob_under * 100.0, 1),
            "badge_color": "#ff3366",
            "valor": True,
        }

    return {
        "pick": "LÍNEA EFICIENTE",
        "linea": t_mkt,
        "cuota": c_over,
        "edge_pts": round(abs(edge_pts), 1),
        "ev_pct": 0.0,
        "prob": round(max(prob_over, prob_under) * 100.0, 1),
        "badge_color": "#555555",
        "valor": False,
    }


def evaluar_valor_spread_nba(
    spread_modelo: float,
    spread_mercado_home: float,
    cuota_home: float,
    cuota_away: float,
    umbral_pts: float = 3.0,
    home_abbr: str = "HOME",
    away_abbr: str = "AWAY",
    prob_home_cover_sim: Optional[float] = None,
    prob_away_cover_sim: Optional[float] = None,
) -> Dict[str, Any]:
    """Auditoría matemática de Hándicap / Spread en formato estrictamente DECIMAL."""
    sp_mod = float(spread_modelo)
    sp_mkt_home = round(float(spread_mercado_home), 1)
    sp_mkt_away = round(-sp_mkt_home, 1)
    c_home = round(float(cuota_home), 2)
    c_away = round(float(cuota_away), 2)

    edge_home_pts = sp_mkt_home - sp_mod
    edge_away_pts = -edge_home_pts

    if prob_home_cover_sim is not None and 0.01 < float(prob_home_cover_sim) < 0.99:
        prob_home_cover = float(prob_home_cover_sim)
    else:
        prob_home_cover = float(np.clip(_normal_cdf(edge_home_pts / NBA_MARGIN_SIM_STD), 0.05, 0.95))

    if prob_away_cover_sim is not None and 0.01 < float(prob_away_cover_sim) < 0.99:
        prob_away_cover = float(prob_away_cover_sim)
    else:
        prob_away_cover = float(np.clip(1.0 - prob_home_cover, 0.05, 0.95))

    if edge_home_pts >= float(umbral_pts):
        ev_pct = round(((prob_home_cover * c_home) - 1.0) * 100.0, 1)
        linea_str = f"{sp_mkt_home:+.1f}"
        return {
            "pick": f"{home_abbr} {linea_str}",
            "lado": "HOME",
            "franquicia": home_abbr,
            "linea": sp_mkt_home,
            "linea_str": linea_str,
            "cuota": c_home,
            "edge_pts": round(edge_home_pts, 1),
            "ev_pct": ev_pct,
            "prob": round(prob_home_cover * 100.0, 1),
            "badge_color": "#00e5ff",
            "valor": True,
        }

    if edge_away_pts >= float(umbral_pts):
        ev_pct = round(((prob_away_cover * c_away) - 1.0) * 100.0, 1)
        linea_str = f"{sp_mkt_away:+.1f}"
        return {
            "pick": f"{away_abbr} {linea_str}",
            "lado": "AWAY",
            "franquicia": away_abbr,
            "linea": sp_mkt_away,
            "linea_str": linea_str,
            "cuota": c_away,
            "edge_pts": round(edge_away_pts, 1),
            "ev_pct": ev_pct,
            "prob": round(prob_away_cover * 100.0, 1),
            "badge_color": "#00e5ff",
            "valor": True,
        }

    return {
        "pick": "SPREAD EFICIENTE",
        "lado": "NEUTRAL",
        "franquicia": home_abbr,
        "linea": sp_mkt_home,
        "linea_str": f"{sp_mkt_home:+.1f}",
        "cuota": c_home,
        "edge_pts": round(abs(edge_home_pts), 1),
        "ev_pct": 0.0,
        "prob": round(max(prob_home_cover, prob_away_cover) * 100.0, 1),
        "badge_color": "#555555",
        "valor": False,
    }


# ==============================================================================
# GESTIÓN DE PLANTILLAS EXCEL (.XLSX): GENERACIÓN Y PARSEO A LINEUPS
# ==============================================================================
def _interpretar_booleano_estrella(val) -> bool:
    """Normaliza valores booleanos o de texto ('True', 'SI', 'SÍ', '1', 'YES') para Es_Estrella."""
    if isinstance(val, (bool, np.bool_)):
        return bool(val)
    if pd.isna(val):
        return False
    s = str(val).strip().upper()
    return s in {"TRUE", "VERDADERO", "1", "1.0", "SI", "SÍ", "YES", "Y", "STAR", "ESTRELLA"}


def generar_dataframe_plantilla_excel(
    lineups_dict: Optional[Dict[str, pd.DataFrame]] = None,
    df_teams: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    """
    Consolida en un solo DataFrame plano la información de los 30 equipos para exportar a Excel (.xlsx).
    Incluye tanto las columnas operativas como las métricas base ('Minutos_Base', 'PPG_Base', 'RPG_Base', 'APG_Base').
    """
    if df_teams is None:
        if (st is not None and hasattr(st, "session_state")) and "df_nba" in st.session_state:
            df_teams = st.session_state.df_nba
        else:
            df_teams = pd.DataFrame(DEFAULT_CSV_DATA)
    if lineups_dict is None:
        if (st is not None and hasattr(st, "session_state")) and "lineups_nba" in st.session_state:
            lineups_dict = st.session_state.lineups_nba
        else:
            lineups_dict = build_default_lineups_dict()

    team_meta_map: Dict[str, Dict] = {}
    for _, row in df_teams.iterrows():
        t_abbr = str(row["Team_Abbr"]).strip().upper()
        team_meta_map[t_abbr] = {
            "Star1_Name": str(row.get("Star1_Name", "")).strip(),
            "Star1_Line": float(row.get("Star1_Line", 20.5)),
            "Star2_Name": str(row.get("Star2_Name", "")).strip(),
            "Star2_Line": float(row.get("Star2_Line", 18.5)),
        }

    records: List[Dict] = []
    ordered_abbrs = list(DEFAULT_LINEUPS_30_TEAMS.keys())

    for abbr in ordered_abbrs:
        lu_df = obtener_lineup_en_sesion(abbr, fallback_lineups=lineups_dict)
        default_lu = {
            str(r["Jugador"]).lower(): r for r in DEFAULT_LINEUPS_30_TEAMS.get(abbr, [])
        }

        meta = team_meta_map.get(abbr, {})
        s1_name = meta.get("Star1_Name", "").lower()
        s2_name = meta.get("Star2_Name", "").lower()
        s1_line = meta.get("Star1_Line", 22.5)
        s2_line = meta.get("Star2_Line", 19.5)

        for _, r in lu_df.iterrows():
            raw_pos = str(r.get("Posición", r.get("Pos", "PG"))).strip()
            pos_label = "Banca" if raw_pos.upper() in {"BENCH", "BANCA"} else raw_pos.upper()
            player_name = str(r.get("Jugador", "")).strip()
            mins_val = round(float(r.get("Minutos_Esperados", r.get("Minutos", 0.0))), 1)
            def_r = default_lu.get(player_name.lower(), {})
            mins_base_val = round(
                float(r.get("Minutos_Base", def_r.get("Minutos_Esperados", mins_val if mins_val > 0 else 34.0))),
                1,
            )
            o_epm_val = round(float(r.get("O_EPM", 0.0)), 1)
            d_epm_val = round(float(r.get("D_EPM", 0.0)), 1)
            pace_diff_val = round(float(r.get("Pace_Diff", 0.0)), 2)

            ppg_def, rpg_def, apg_def = _derivar_metricas_base_jugador(
                player_name, pos_label, mins_base_val, o_epm_val, d_epm_val, abbr
            )
            ppg_base_val = round(float(r.get("PPG_Base", ppg_def)), 1)
            rpg_base_val = round(float(r.get("RPG_Base", rpg_def)), 1)
            apg_base_val = round(float(r.get("APG_Base", apg_def)), 1)

            if "Es_Estrella" in r and pd.notna(r["Es_Estrella"]):
                es_estrella = _interpretar_booleano_estrella(r["Es_Estrella"])
            else:
                p_low = player_name.lower()
                es_estrella = bool(pos_label != "Banca" and p_low in {s1_name, s2_name})

            if "PPG_Linea" in r and pd.notna(r["PPG_Linea"]) and float(r["PPG_Linea"]) > 0:
                ppg_linea = round(float(r["PPG_Linea"]), 1)
            elif player_name.lower() == s1_name:
                ppg_linea = round(float(s1_line), 1)
            elif player_name.lower() == s2_name:
                ppg_linea = round(float(s2_line), 1)
            elif pos_label == "Banca":
                ppg_linea = 0.0
            else:
                ppg_linea = round_to_half_point(max(0.5, ppg_base_val - 0.5))

            records.append(
                {
                    "Team_Abbr": abbr,
                    "Pos": pos_label,
                    "Jugador": player_name,
                    "Minutos": float(mins_val),
                    "Minutos_Base": float(mins_base_val),
                    "O_EPM": float(o_epm_val),
                    "D_EPM": float(d_epm_val),
                    "Pace_Diff": float(pace_diff_val),
                    "Es_Estrella": bool(es_estrella),
                    "PPG_Base": float(ppg_base_val),
                    "RPG_Base": float(rpg_base_val),
                    "APG_Base": float(apg_base_val),
                    "PPG_Linea": float(ppg_linea),
                }
            )

    df_out = pd.DataFrame(records)[EXCEL_TEMPLATE_COLS]
    for fcol in ["Minutos", "Minutos_Base", "O_EPM", "D_EPM", "Pace_Diff", "PPG_Base", "RPG_Base", "APG_Base", "PPG_Linea"]:
        df_out[fcol] = df_out[fcol].astype(float)
    return df_out


def parsear_excel_a_lineups(df_excel: pd.DataFrame, strict_240: bool = False) -> Dict[str, pd.DataFrame]:
    """
    Valida y parsea el DataFrame cargado desde Excel (.xlsx).
    Soporta o deriva automáticamente las métricas base por jugador:
      - 'PPG_Base', 'RPG_Base', 'APG_Base' y 'Minutos_Base'.
    """
    if df_excel is None or not isinstance(df_excel, pd.DataFrame) or df_excel.empty:
        raise ValueError("El archivo Excel subido está vacío o no contiene una tabla válida.")

    df_clean = df_excel.copy()
    df_clean.columns = [str(c).strip() for c in df_clean.columns]

    missing_cols = [c for c in EXCEL_TEMPLATE_BASE_COLS if c not in df_clean.columns]
    if missing_cols:
        raise ValueError(
            f"El archivo Excel no contiene las columnas requeridas: {', '.join(missing_cols)}. "
            f"Columnas base esperadas: {', '.join(EXCEL_TEMPLATE_BASE_COLS)}"
        )

    df_clean = df_clean.dropna(subset=["Team_Abbr", "Jugador"], how="all").copy()
    df_clean["Team_Abbr"] = df_clean["Team_Abbr"].astype(str).str.strip().str.upper()
    df_clean["Minutos"] = pd.to_numeric(df_clean["Minutos"], errors="coerce").fillna(0.0).clip(lower=0.0).astype(float)
    df_clean["O_EPM"] = pd.to_numeric(df_clean["O_EPM"], errors="coerce").fillna(0.0).astype(float)
    df_clean["D_EPM"] = pd.to_numeric(df_clean["D_EPM"], errors="coerce").fillna(0.0).astype(float)
    df_clean["Pace_Diff"] = pd.to_numeric(df_clean["Pace_Diff"], errors="coerce").fillna(0.0).astype(float)
    df_clean["PPG_Linea"] = pd.to_numeric(df_clean["PPG_Linea"], errors="coerce").fillna(0.0).astype(float)
    df_clean["Es_Estrella"] = df_clean["Es_Estrella"].apply(_interpretar_booleano_estrella)

    has_min_base = "Minutos_Base" in df_clean.columns
    has_ppg_base = "PPG_Base" in df_clean.columns
    has_rpg_base = "RPG_Base" in df_clean.columns
    has_apg_base = "APG_Base" in df_clean.columns

    lineups_parsed: Dict[str, pd.DataFrame] = build_default_lineups_dict()

    for team_abbr, group in df_clean.groupby("Team_Abbr", sort=False):
        total_team_mins = float(group["Minutos"].sum())
        if strict_240 and abs(total_team_mins - 240.0) > 0.5:
            raise ValueError(
                f"Los minutos del equipo '{team_abbr}' suman {total_team_mins:.1f} "
                f"(deben sumar exactamente 240.0 minutos por franquicia)."
            )

        default_lu = {
            str(r["Jugador"]).lower(): r for r in DEFAULT_LINEUPS_30_TEAMS.get(team_abbr, [])
        }

        team_rows: List[Dict] = []
        for _, row in group.iterrows():
            raw_pos = str(row["Pos"]).strip().upper()
            internal_pos = "BENCH" if raw_pos in {"BANCA", "BENCH", "BN", "ROTACION", "ROTACIÓN"} else raw_pos
            p_name = str(row["Jugador"]).strip()
            mins_val = round(float(row["Minutos"]), 1)
            o_epm_val = round(float(row["O_EPM"]), 1)
            d_epm_val = round(float(row["D_EPM"]), 1)
            def_r = default_lu.get(p_name.lower(), {})
            def_mins = float(def_r.get("Minutos_Esperados", mins_val if mins_val > 0 else 34.0))

            if has_min_base and pd.notna(row.get("Minutos_Base")) and float(row["Minutos_Base"]) > 0:
                mins_base_val = round(float(row["Minutos_Base"]), 1)
            else:
                mins_base_val = round(def_mins if def_mins > 0 else 34.0, 1)

            ppg_der, rpg_der, apg_der = _derivar_metricas_base_jugador(
                p_name, internal_pos, mins_base_val, o_epm_val, d_epm_val, team_abbr
            )

            ppg_base_val = (
                round(float(row["PPG_Base"]), 1)
                if has_ppg_base and pd.notna(row.get("PPG_Base")) and float(row["PPG_Base"]) > 0
                else ppg_der
            )
            rpg_base_val = (
                round(float(row["RPG_Base"]), 1)
                if has_rpg_base and pd.notna(row.get("RPG_Base")) and float(row["RPG_Base"]) > 0
                else rpg_der
            )
            apg_base_val = (
                round(float(row["APG_Base"]), 1)
                if has_apg_base and pd.notna(row.get("APG_Base")) and float(row["APG_Base"]) > 0
                else apg_der
            )

            team_rows.append(
                {
                    "Posición": internal_pos,
                    "Jugador": p_name,
                    "Minutos_Esperados": mins_val,
                    "Minutos_Base": mins_base_val,
                    "O_EPM": o_epm_val,
                    "D_EPM": d_epm_val,
                    "Pace_Diff": round(float(row["Pace_Diff"]), 2),
                    "Es_Estrella": bool(row["Es_Estrella"]),
                    "PPG_Base": ppg_base_val,
                    "RPG_Base": rpg_base_val,
                    "APG_Base": apg_base_val,
                    "PPG_Linea": round(float(row["PPG_Linea"]), 1),
                }
            )

        lineups_parsed[team_abbr] = pd.DataFrame(team_rows)

    return lineups_parsed


def construir_excel_plantilla_bytes(
    lineups_dict: Optional[Dict[str, pd.DataFrame]] = None,
    df_teams: Optional[pd.DataFrame] = None,
    historial_calibracion: Optional[List[Dict[str, Any]]] = None,
) -> bytes:
    """
    Convierte el DataFrame de `generar_dataframe_plantilla_excel()` a formato Excel binario (.xlsx)
    en memoria usando `io.BytesIO()` y `openpyxl`, incluyendo:
      - Hoja 1 (`Rosters_NBA_30`): Plantilla plana de rosters y métricas individuales calibradas.
      - Hoja 2 (`Ratings_Equipos_Calibrados`): Tabla de las 30 franquicias con Pace, Off_Rtg y Def_Rtg actuales.
      - Hoja 3 (`Bitacora_Auditoria_Kalman`, opcional si existen calibraciones): Historial de residuales post-partido.
    """
    df_plantilla = generar_dataframe_plantilla_excel(lineups_dict=lineups_dict, df_teams=df_teams)
    if df_teams is None:
        if (st is not None and hasattr(st, "session_state")) and "df_nba" in st.session_state:
            df_teams = st.session_state.df_nba
        else:
            df_teams = pd.DataFrame(DEFAULT_CSV_DATA)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        # Hoja 1: Rosters_NBA_30 (compatible directamente con st.file_uploader)
        df_plantilla.to_excel(writer, index=False, sheet_name="Rosters_NBA_30")
        ws = writer.sheets["Rosters_NBA_30"]
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        header_fill = PatternFill(start_color="0B1223", end_color="0B1223", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="00F5D4")
        alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for col_idx in range(1, len(EXCEL_TEMPLATE_COLS) + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center" if col_idx != 3 else "left", vertical="center")
        ws.row_dimensions[1].height = 24

        for row_idx in range(2, len(df_plantilla) + 2):
            team_block = (row_idx - 2) // 6
            for col_idx in range(1, len(EXCEL_TEMPLATE_COLS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if team_block % 2 == 1:
                    cell.fill = alt_fill
                if col_idx in (1, 2, 9):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 3:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.number_format = "0.0"

        col_widths = {
            "A": 12, "B": 10, "C": 30, "D": 11, "E": 14,
            "F": 10, "G": 10, "H": 11, "I": 13, "J": 12, "K": 12, "L": 12, "M": 12,
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        # Hoja 2: Ratings_Equipos_Calibrados
        team_export_cols = [
            c for c in ["Team_Abbr", "Team", "Conference", "Division", "Pace", "Off_Rtg", "Def_Rtg", "Star1_Name", "Star1_PPG", "Star2_Name", "Star2_PPG"]
            if c in df_teams.columns
        ]
        df_teams[team_export_cols].to_excel(writer, index=False, sheet_name="Ratings_Equipos_Calibrados")
        ws_t = writer.sheets["Ratings_Equipos_Calibrados"]
        ws_t.views.sheetView[0].showGridLines = True
        ws_t.freeze_panes = "A2"
        for col_idx in range(1, len(team_export_cols) + 1):
            c = ws_t.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font

        # Hoja 3: Bitacora_Auditoria_Kalman (si hay registros)
        if historial_calibracion:
            df_hist = pd.DataFrame(historial_calibracion)
            df_hist.to_excel(writer, index=False, sheet_name="Bitacora_Auditoria_Kalman")
            ws_h = writer.sheets["Bitacora_Auditoria_Kalman"]
            ws_h.views.sheetView[0].showGridLines = True
            ws_h.freeze_panes = "A2"
            for col_idx in range(1, len(df_hist.columns) + 1):
                c = ws_h.cell(row=1, column=col_idx)
                c.fill = header_fill
                c.font = header_font

    buffer.seek(0)
    return buffer.getvalue()


# ==============================================================================
# MOTOR DE CÁLCULO ASCENDENTE (BOTTOM-UP LINEUP AGGREGATOR)
# ==============================================================================
def calcular_ratings_lineup(df_lineup: pd.DataFrame, pace_base: float) -> Tuple[float, float, float]:
    """
    Proyecta el Off_Rtg, Def_Rtg y Pace sumando los minutos ponderados y el impacto
    neto individual (O_EPM y D_EPM) de los 5 titulares más la unidad de banca (240 minutos totales).
    """
    df_clean = df_lineup.copy()
    mins_col = "Minutos_Esperados" if "Minutos_Esperados" in df_clean.columns else "Minutos"
    pos_col = "Posición" if "Posición" in df_clean.columns else "Pos"

    mins = pd.to_numeric(df_clean[mins_col], errors="coerce").fillna(0.0).clip(lower=0.0).to_numpy(dtype=float)
    o_epm = pd.to_numeric(df_clean["O_EPM"], errors="coerce").fillna(0.0).to_numpy(dtype=float)
    d_epm = pd.to_numeric(df_clean["D_EPM"], errors="coerce").fillna(0.0).to_numpy(dtype=float)
    pos_arr = df_clean[pos_col].astype(str).str.upper().to_numpy()
    pace_diff_arr = (
        pd.to_numeric(df_clean["Pace_Diff"], errors="coerce").fillna(0.0).to_numpy(dtype=float)
        if "Pace_Diff" in df_clean.columns
        else np.zeros(len(df_clean), dtype=float)
    )

    total_mins = float(np.sum(mins))
    missing_mins = 0.0

    if total_mins <= 0.0:
        eff_mins = np.array([32.0, 32.0, 32.0, 32.0, 32.0, 80.0][: len(mins)], dtype=float)
    elif total_mins < 240.0:
        missing_mins = 240.0 - total_mins
        eff_mins = mins.copy()
    elif total_mins > 240.0:
        eff_mins = mins * (240.0 / total_mins)
    else:
        eff_mins = mins.copy()

    court_slots = eff_mins / 48.0
    replacement_slots = missing_mins / 48.0

    synergy_scale = 0.62
    team_off_delta = float((np.sum(court_slots * o_epm) + replacement_slots * (-2.2)) * synergy_scale)
    team_def_delta = float((np.sum(court_slots * d_epm) + replacement_slots * (-1.2)) * synergy_scale)

    proj_off_rtg = float(np.clip(LEAGUE_AVG_ORTG + team_off_delta, 98.0, 132.0))
    proj_def_rtg = float(np.clip(LEAGUE_AVG_ORTG - team_def_delta, 98.0, 132.0))

    guard_mask = np.isin(pos_arr, ["PG", "SG"])
    big_mask = np.isin(pos_arr, ["PF", "C"])
    guard_mins = float(np.sum(eff_mins[guard_mask])) if np.any(guard_mask) else 64.0
    guard_o_epm = float(np.mean(o_epm[guard_mask])) if np.any(guard_mask) else 0.0
    big_d_epm = float(np.mean(d_epm[big_mask])) if np.any(big_mask) else 0.0

    pace_guard_boost = ((guard_mins - 65.0) / 12.0) * 0.45
    pace_style_boost = (guard_o_epm * 0.16) - (big_d_epm * 0.10)
    custom_pace_diff = float(np.sum((eff_mins / 48.0) * pace_diff_arr))
    proj_pace = float(np.clip(float(pace_base) + pace_guard_boost + pace_style_boost + custom_pace_diff, 88.0, 108.0))

    return round(proj_off_rtg, 1), round(proj_def_rtg, 1), round(proj_pace, 1)


def apply_lineup_to_team_row(team_row: pd.Series, df_lineup: pd.DataFrame) -> pd.Series:
    """
    Actualiza dinámicamente Off_Rtg, Def_Rtg, Pace y los 2 líderes anotadores (Star1 / Star2)
    de la fila del equipo a partir de su quinteto en sesión.
    """
    updated = team_row.copy()
    t_abbr = str(team_row["Team_Abbr"]).strip().upper()
    default_meta = DEFAULT_TEAM_META_BY_ABBR.get(t_abbr, {})
    default_lu = {
        str(r["Jugador"]).lower(): r for r in DEFAULT_LINEUPS_30_TEAMS.get(t_abbr, [])
    }

    pace_base = float(team_row.get("Base_Pace", default_meta.get("Pace", team_row["Pace"])))
    base_off = float(team_row.get("Base_Off_Rtg", default_meta.get("Off_Rtg", team_row["Off_Rtg"])))
    base_def = float(team_row.get("Base_Def_Rtg", default_meta.get("Def_Rtg", team_row["Def_Rtg"])))

    off_rtg, def_rtg, dyn_pace = calcular_ratings_lineup(df_lineup, pace_base)

    updated["Base_Off_Rtg"] = base_off
    updated["Base_Def_Rtg"] = base_def
    updated["Base_Pace"] = pace_base

    updated["Off_Rtg"] = off_rtg
    updated["Def_Rtg"] = def_rtg
    updated["Pace"] = dyn_pace

    pos_col = "Posición" if "Posición" in df_lineup.columns else "Pos"
    mins_col = "Minutos_Esperados" if "Minutos_Esperados" in df_lineup.columns else "Minutos"

    starters_df = df_lineup[~df_lineup[pos_col].astype(str).str.upper().isin(["BENCH", "BANCA"])].copy()
    if starters_df.empty:
        return updated

    starters_df["Minutos_Esperados"] = pd.to_numeric(starters_df[mins_col], errors="coerce").fillna(0.0)
    starters_df["O_EPM"] = pd.to_numeric(starters_df["O_EPM"], errors="coerce").fillna(0.0)
    starters_df["score_impacto"] = starters_df["Minutos_Esperados"] * (1.0 + 0.16 * starters_df["O_EPM"])

    def _calcular_ppg_y_linea_jugador(st_row: pd.Series, star_slot: int) -> Tuple[str, float, float, float, float, float, float]:
        p_name = str(st_row["Jugador"]).strip()
        p_pos = str(st_row.get(pos_col, "SF")).strip().upper()
        p_mins = float(st_row["Minutos_Esperados"])
        p_oepm = float(st_row["O_EPM"])
        p_depm = float(st_row.get("D_EPM", 0.0))
        excel_line = float(st_row.get("PPG_Linea", 0.0)) if "PPG_Linea" in st_row and pd.notna(st_row["PPG_Linea"]) else 0.0

        orig_s1 = str(default_meta.get("Star1_Name", "")).strip()
        orig_s2 = str(default_meta.get("Star2_Name", "")).strip()
        def_row = default_lu.get(p_name.lower(), {})
        def_mins = float(st_row.get("Minutos_Base", def_row.get("Minutos_Esperados", 35.0 if star_slot == 1 else 33.0)))
        if def_mins <= 0.5:
            def_mins = 34.0
        def_oepm = float(def_row.get("O_EPM", p_oepm))

        ppg_der, rpg_der, apg_der = _derivar_metricas_base_jugador(
            p_name, p_pos, def_mins, p_oepm, p_depm, t_abbr
        )

        if "PPG_Base" in st_row and pd.notna(st_row["PPG_Base"]) and float(st_row["PPG_Base"]) > 0:
            orig_ppg = float(st_row["PPG_Base"])
        elif p_name.lower() == orig_s1.lower():
            orig_ppg = float(default_meta.get("Star1_PPG", 25.0))
        elif p_name.lower() == orig_s2.lower():
            orig_ppg = float(default_meta.get("Star2_PPG", 21.0))
        else:
            orig_ppg = ppg_der

        rpg_base = float(st_row["RPG_Base"]) if "RPG_Base" in st_row and pd.notna(st_row["RPG_Base"]) and float(st_row["RPG_Base"]) > 0 else rpg_der
        apg_base = float(st_row["APG_Base"]) if "APG_Base" in st_row and pd.notna(st_row["APG_Base"]) and float(st_row["APG_Base"]) > 0 else apg_der

        if p_name.lower() == orig_s1.lower():
            orig_line = float(default_meta.get("Star1_Line", 24.5))
        elif p_name.lower() == orig_s2.lower():
            orig_line = float(default_meta.get("Star2_Line", 20.5))
        else:
            orig_line = excel_line if excel_line > 0 else round_to_half_point(orig_ppg - 0.5)

        mins_changed = abs(p_mins - def_mins) > 0.2
        epm_changed = abs(p_oepm - def_oepm) > 0.05
        line_changed = excel_line > 0.0 and abs(excel_line - orig_line) > 0.1

        proj_ppg = orig_ppg * (p_mins / max(def_mins, 1.0)) + (p_oepm - def_oepm) * 1.15
        if line_changed and not mins_changed and not epm_changed:
            prop_line = round_to_half_point(excel_line)
        elif excel_line > 0.0:
            prop_line = round_to_half_point(excel_line)
        else:
            prop_line = round_to_half_point(max(0.5, proj_ppg - 0.5))

        return (
            p_name,
            round(max(4.0, proj_ppg), 1),
            prop_line,
            round(orig_ppg, 1),
            round(rpg_base, 1),
            round(apg_base, 1),
            round(def_mins, 1),
        )

    active_starters = starters_df[starters_df["Minutos_Esperados"] > 0.5].copy()
    if active_starters.empty:
        active_starters = starters_df.copy()

    if "Es_Estrella" in active_starters.columns:
        active_starters["_star_flag"] = active_starters["Es_Estrella"].apply(_interpretar_booleano_estrella)
        active_starters = active_starters.sort_values(
            by=["_star_flag", "score_impacto"], ascending=[False, False]
        )
    else:
        orig_names = [
            str(updated.get("Star1_Name", "")).lower(),
            str(updated.get("Star2_Name", "")).lower(),
        ]
        active_starters["_star_flag"] = active_starters["Jugador"].astype(str).str.lower().isin(orig_names)
        active_starters = active_starters.sort_values(
            by=["_star_flag", "score_impacto"], ascending=[False, False]
        )

    for star_slot in (1, 2):
        if len(active_starters) >= star_slot:
            chosen_row = active_starters.iloc[star_slot - 1]
            s_name, s_ppg, s_line, s_ppg_base, s_rpg_base, s_apg_base, s_min_base = _calcular_ppg_y_linea_jugador(
                chosen_row, star_slot
            )
            updated[f"Star{star_slot}_Name"] = s_name
            updated[f"Star{star_slot}_PPG"] = s_ppg
            updated[f"Star{star_slot}_Line"] = s_line
            updated[f"Star{star_slot}_PPG_Base"] = s_ppg_base
            updated[f"Star{star_slot}_RPG_Base"] = s_rpg_base
            updated[f"Star{star_slot}_APG_Base"] = s_apg_base
            updated[f"Star{star_slot}_Minutos_Base"] = s_min_base
            updated[f"Star{star_slot}_Minutos_Proy"] = round(float(chosen_row["Minutos_Esperados"]), 1)

    return updated


def sincronizar_df_nba_con_lineups(
    df_nba: pd.DataFrame, lineups_dict: Dict[str, pd.DataFrame]
) -> pd.DataFrame:
    """
    Sincroniza inmediatamente `df_nba` (usado en el Gestor de Rosters y cabeceras)
    con los quintetos cargados en `st.session_state.lineups_nba`.
    """
    df_sync = df_nba.copy()
    for idx, row in df_sync.iterrows():
        abbr = str(row["Team_Abbr"]).strip().upper()
        lu_df = obtener_lineup_en_sesion(abbr, fallback_lineups=lineups_dict)
        updated_row = apply_lineup_to_team_row(row, lu_df)
        for col in [
            "Pace", "Off_Rtg", "Def_Rtg",
            "Base_Pace", "Base_Off_Rtg", "Base_Def_Rtg",
            "Star1_Name", "Star1_PPG", "Star1_Line",
            "Star1_PPG_Base", "Star1_RPG_Base", "Star1_APG_Base", "Star1_Minutos_Base", "Star1_Minutos_Proy",
            "Star2_Name", "Star2_PPG", "Star2_Line",
            "Star2_PPG_Base", "Star2_RPG_Base", "Star2_APG_Base", "Star2_Minutos_Base", "Star2_Minutos_Proy",
        ]:
            if col in updated_row:
                df_sync.loc[idx, col] = updated_row[col]
    return df_sync


def simulate_player_prop(
    rng: np.random.Generator,
    player_name: str,
    team_abbr: str,
    opp_abbr: str,
    base_ppg: float,
    prop_line: float,
    style: str,
    player_std: float,
    team_pace: float,
    exp_pace: float,
    team_style_off_ppp: float,
    opp_style_def_ppp: float,
    opp_def_rtg: float,
    poss_sim: np.ndarray,
    ppg_base_raw: Optional[float] = None,
    rpg_base: float = 5.5,
    apg_base: float = 4.8,
    minutos_base: float = 34.0,
    minutos_proyectados: float = 34.0,
) -> Dict:
    """
    Motor combinado para Player Props:
    - Simula la varianza de puntos e integra `proyectar_props_jugador` para Puntos (PTS),
      Rebotes (REB) y Asistencias (AST) con auditoría Poisson +EV en cuotas decimales.
    """
    n_sims = len(poss_sim)
    pace_factor = exp_pace / max(team_pace, 85.0)

    league_style_avg = LEAGUE_AVG_ISO_PPP if style.upper() == "ISO" else LEAGUE_AVG_PNR_PPP
    style_def_ratio = opp_style_def_ppp / league_style_avg
    off_synergy_ratio = team_style_off_ppp / league_style_avg
    style_multiplier = 1.0 + 0.58 * (style_def_ratio - 1.0) + 0.18 * (off_synergy_ratio - 1.0)

    team_def_multiplier = 1.0 + 0.30 * ((opp_def_rtg - LEAGUE_AVG_ORTG) / LEAGUE_AVG_ORTG)
    adjusted_ppg = base_ppg * pace_factor * style_multiplier * team_def_multiplier

    iter_pace_scale = poss_sim / exp_pace
    iter_mean = adjusted_ppg * iter_pace_scale

    eff_std = player_std * np.sqrt(pace_factor)
    sim_pts = np.clip(rng.normal(loc=iter_mean, scale=eff_std, size=n_sims), 0.0, None)

    prob_over = float(np.mean(sim_pts > prop_line))
    prob_under = 1.0 - prob_over
    lean_side = "OVER" if prob_over >= prob_under else "UNDER"
    lean_prob = prob_over if prob_over >= prob_under else prob_under

    dec_odds, _ = prob_to_fair_odds(lean_prob)
    style_label = "Iso_PPP_Def" if style.upper() == "ISO" else "PnR_PPP_Def"

    jugador_dict = {
        "Jugador": player_name,
        "Team_Abbr": team_abbr,
        "PPG_Base": float(ppg_base_raw if ppg_base_raw is not None else base_ppg),
        "RPG_Base": float(rpg_base),
        "APG_Base": float(apg_base),
        "Minutos_Base": float(minutos_base),
        "PPG_Linea": float(prop_line),
        "RPG_Linea": round_to_half_point(max(0.5, rpg_base - 0.5)),
        "APG_Linea": round_to_half_point(max(0.5, apg_base - 0.5)),
    }
    multi_props = proyectar_props_jugador(
        jugador_dict=jugador_dict,
        minutos_proyectados=minutos_proyectados,
        pace_partido=exp_pace,
        pace_promedio_liga=99.5,
        def_rating_rival=opp_def_rtg,
    )

    return {
        "player_name": player_name,
        "team_abbr": team_abbr,
        "opp_abbr": opp_abbr,
        "style": style.upper(),
        "base_ppg": float(base_ppg),
        "ppg_base_raw": float(jugador_dict["PPG_Base"]),
        "rpg_base": float(rpg_base),
        "apg_base": float(apg_base),
        "minutos_base": float(minutos_base),
        "minutos_proyectados": float(minutos_proyectados),
        "opp_def_rtg": float(opp_def_rtg),
        "exp_pace": float(exp_pace),
        "proj_ppg": float(np.mean(sim_pts)),
        "prop_line": float(prop_line),
        "prob_over": prob_over,
        "prob_under": prob_under,
        "lean_side": lean_side,
        "lean_prob": lean_prob,
        "fair_dec": dec_odds,
        "fair_amer": dec_odds,
        "edge_pts": float(adjusted_ppg - prop_line),
        "p25": float(np.percentile(sim_pts, 25)),
        "p75": float(np.percentile(sim_pts, 75)),
        "matchup_note": f"vs {opp_abbr} {style_label} ({opp_style_def_ppp:.2f} PPP) | DefRtg {opp_def_rtg:.1f}",
        "jugador_dict": jugador_dict,
        "multi_props": multi_props,
    }


def run_monte_carlo_matchup(
    away_row_raw: pd.Series,
    home_row_raw: pd.Series,
    away_lineup: Optional[pd.DataFrame] = None,
    home_lineup: Optional[pd.DataFrame] = None,
    n_sims: int = 5000,
    hfa_pts: float = 2.5,
    custom_spread_home: Optional[float] = None,
    custom_total_line: Optional[float] = None,
    market_odds_info: Optional[Dict[str, Any]] = None,
    umbral_ev_pts: float = 3.0,
    seed: int = 42,
) -> Dict:
    """
    Integra el Motor de Cálculo Ascendente (`calcular_ratings_lineup`) leyendo obligatoriamente
    los quintetos en sesión (`st.session_state.lineups_nba`), ejecuta 5,000 simulaciones de Monte Carlo
    y audita tanto las líneas de partido como los Player Props (PTS, REB, AST) en FORMATO DECIMAL.
    """
    away_abbr = str(away_row_raw["Team_Abbr"]).strip().upper()
    home_abbr = str(home_row_raw["Team_Abbr"]).strip().upper()

    if away_lineup is None:
        away_lineup = obtener_lineup_en_sesion(away_abbr)
    if home_lineup is None:
        home_lineup = obtener_lineup_en_sesion(home_abbr)

    away_row = apply_lineup_to_team_row(away_row_raw, away_lineup)
    home_row = apply_lineup_to_team_row(home_row_raw, home_lineup)

    rng = np.random.default_rng(seed)

    away_pace = float(away_row["Pace"])
    home_pace = float(home_row["Pace"])
    exp_possessions = (away_pace + home_pace) / 2.0

    away_base_ortg = float(away_row["Off_Rtg"]) + 0.58 * (float(home_row["Def_Rtg"]) - LEAGUE_AVG_ORTG)
    home_base_ortg = float(home_row["Off_Rtg"]) + 0.58 * (float(away_row["Def_Rtg"]) - LEAGUE_AVG_ORTG)

    static_away_ortg = float(away_row["Base_Off_Rtg"]) + 0.58 * (float(home_row["Base_Def_Rtg"]) - LEAGUE_AVG_ORTG)
    static_home_ortg = float(home_row["Base_Off_Rtg"]) + 0.58 * (float(away_row["Base_Def_Rtg"]) - LEAGUE_AVG_ORTG)
    static_pace = (float(away_row["Base_Pace"]) + float(home_row["Base_Pace"])) / 2.0
    mkt_away_pts = (static_pace / 100.0) * static_away_ortg - (hfa_pts / 2.0)
    mkt_home_pts = (static_pace / 100.0) * static_home_ortg + (hfa_pts / 2.0)
    default_market_spread_home = float(np.floor(-(mkt_home_pts - mkt_away_pts)) + 0.5)
    default_market_total = float(np.floor(mkt_home_pts + mkt_away_pts) + 0.5)

    away_iso_delta = float(away_row["Iso_PPP_Off"]) - float(home_row["Iso_PPP_Def"])
    away_pnr_delta = float(away_row["PnR_PPP_Off"]) - float(home_row["PnR_PPP_Def"])
    home_iso_delta = float(home_row["Iso_PPP_Off"]) - float(away_row["Iso_PPP_Def"])
    home_pnr_delta = float(home_row["PnR_PPP_Off"]) - float(away_row["PnR_PPP_Def"])

    away_tactical_ortg_adj = (away_iso_delta * 5.2) + (away_pnr_delta * 7.4)
    home_tactical_ortg_adj = (home_iso_delta * 5.2) + (home_pnr_delta * 7.4)

    away_tov_adj = (LEAGUE_AVG_TOV - float(away_row["TOV_Pct"])) * 0.35
    home_tov_adj = (LEAGUE_AVG_TOV - float(home_row["TOV_Pct"])) * 0.35

    away_adj_ortg = away_base_ortg + away_tactical_ortg_adj + away_tov_adj
    home_adj_ortg = home_base_ortg + home_tactical_ortg_adj + home_tov_adj

    mkt_info = market_odds_info or {}
    mkt_spread_val = mkt_info.get("spread_mercado_home")
    mkt_total_val = mkt_info.get("total_mercado")

    if custom_spread_home is not None:
        spread_line_home = float(custom_spread_home)
    elif mkt_spread_val is not None:
        spread_line_home = float(mkt_spread_val)
    else:
        spread_line_home = default_market_spread_home

    if custom_total_line is not None:
        total_line = float(custom_total_line)
    elif mkt_total_val is not None:
        total_line = float(mkt_total_val)
    else:
        total_line = default_market_total

    spread_line_away = -spread_line_home

    cuota_over_decimal = round(float(mkt_info.get("cuota_over_decimal", 1.91)), 2)
    cuota_under_decimal = round(float(mkt_info.get("cuota_under_decimal", 1.91)), 2)
    cuota_spread_home_decimal = round(float(mkt_info.get("cuota_spread_home_decimal", 1.91)), 2)
    cuota_spread_away_decimal = round(float(mkt_info.get("cuota_spread_away_decimal", 1.91)), 2)
    odds_source_label = str(mkt_info.get("fuente", "Consenso Sportsbook (Decimal)"))

    poss_sim = np.clip(rng.normal(loc=exp_possessions, scale=2.75, size=n_sims), 86.0, 114.0)

    away_ortg_std = 9.2 * (1.0 + 0.35 * (float(away_row["3PA_Rate"]) - 0.38))
    home_ortg_std = 9.2 * (1.0 + 0.35 * (float(home_row["3PA_Rate"]) - 0.38))

    away_ortg_sim = rng.normal(loc=away_adj_ortg, scale=away_ortg_std, size=n_sims)
    home_ortg_sim = rng.normal(loc=home_adj_ortg, scale=home_ortg_std, size=n_sims)

    away_pts_sim = (poss_sim / 100.0) * away_ortg_sim - (hfa_pts / 2.0)
    home_pts_sim = (poss_sim / 100.0) * home_ortg_sim + (hfa_pts / 2.0)

    away_3pa_sim = np.clip(rng.normal(poss_sim * float(away_row["3PA_Rate"]) * 0.92, 4.0, size=n_sims), 18, 60).astype(int)
    home_3pa_sim = np.clip(rng.normal(poss_sim * float(home_row["3PA_Rate"]) * 0.92, 4.0, size=n_sims), 18, 60).astype(int)

    away_3p_pct_base = float(away_row.get("3P_Pct", 0.368))
    home_3p_pct_base = float(home_row.get("3P_Pct", 0.368))

    away_z = (away_ortg_sim - away_adj_ortg) / away_ortg_std
    home_z = (home_ortg_sim - home_adj_ortg) / home_ortg_std

    away_3pm_sim = rng.binomial(n=away_3pa_sim, p=np.clip(away_3p_pct_base + 0.032 * away_z, 0.18, 0.58))
    home_3pm_sim = rng.binomial(n=home_3pa_sim, p=np.clip(home_3p_pct_base + 0.032 * home_z, 0.18, 0.58))

    away_tov_sim = np.clip(rng.normal(loc=(poss_sim / 100.0) * float(away_row["TOV_Pct"]) - 0.85 * away_z, scale=2.2, size=n_sims), 4.0, 28.0)
    home_tov_sim = np.clip(rng.normal(loc=(poss_sim / 100.0) * float(home_row["TOV_Pct"]) - 0.85 * home_z, scale=2.2, size=n_sims), 4.0, 28.0)

    margin_sim = home_pts_sim - away_pts_sim
    ties_mask = np.abs(margin_sim) < 0.5
    if np.any(ties_mask):
        ot_diff = rng.normal(loc=0.35, scale=4.2, size=int(np.sum(ties_mask)))
        ot_diff = np.where(ot_diff >= 0, np.maximum(ot_diff, 1.0), np.minimum(ot_diff, -1.0))
        home_pts_sim[ties_mask] += np.maximum(ot_diff, 0) + 8.0
        away_pts_sim[ties_mask] += np.maximum(-ot_diff, 0) + 8.0
        margin_sim = home_pts_sim - away_pts_sim

    total_pts_sim = home_pts_sim + away_pts_sim

    exp_away_pts = float(np.mean(away_pts_sim))
    exp_home_pts = float(np.mean(home_pts_sim))
    exp_margin_home = float(np.mean(margin_sim))
    exp_total_pts = float(np.mean(total_pts_sim))
    proj_spread_home = -exp_margin_home

    home_win_prob = float(np.mean(margin_sim > 0))
    away_win_prob = 1.0 - home_win_prob
    home_ml_dec, home_ml_amer = prob_to_fair_odds(home_win_prob)
    away_ml_dec, away_ml_amer = prob_to_fair_odds(away_win_prob)

    home_cover_prob = float(np.mean((home_pts_sim + spread_line_home) > away_pts_sim))
    away_cover_prob = 1.0 - home_cover_prob
    home_sp_dec, home_sp_amer = prob_to_fair_odds(home_cover_prob)
    away_sp_dec, away_sp_amer = prob_to_fair_odds(away_cover_prob)

    over_prob = float(np.mean(total_pts_sim > total_line))
    under_prob = 1.0 - over_prob
    over_dec, over_amer = prob_to_fair_odds(over_prob)
    under_dec, under_amer = prob_to_fair_odds(under_prob)

    valor_totales = evaluar_valor_totales_nba(
        total_modelo=exp_total_pts,
        total_mercado=total_line,
        cuota_over=cuota_over_decimal,
        cuota_under=cuota_under_decimal,
        umbral_pts=umbral_ev_pts,
        prob_over_sim=over_prob,
        prob_under_sim=under_prob,
    )
    valor_spread = evaluar_valor_spread_nba(
        spread_modelo=proj_spread_home,
        spread_mercado_home=spread_line_home,
        cuota_home=cuota_spread_home_decimal,
        cuota_away=cuota_spread_away_decimal,
        umbral_pts=umbral_ev_pts,
        home_abbr=home_abbr,
        away_abbr=away_abbr,
        prob_home_cover_sim=home_cover_prob,
        prob_away_cover_sim=away_cover_prob,
    )

    fav_is_home = home_win_prob >= away_win_prob
    fav_abbr = home_abbr if fav_is_home else away_abbr
    fav_name = str(home_row["Team"]) if fav_is_home else str(away_row["Team"])
    fav_base_win = home_win_prob if fav_is_home else away_win_prob
    fav_win_mask = (margin_sim > 0) if fav_is_home else (margin_sim < 0)
    fav_3pm_arr = home_3pm_sim if fav_is_home else away_3pm_sim
    fav_tov_arr = home_tov_sim if fav_is_home else away_tov_sim
    fav_ortg_arr = home_ortg_sim if fav_is_home else away_ortg_sim

    threepm_threshold = int(np.round(np.percentile(fav_3pm_arr, 52)))
    mask_3pm = fav_3pm_arr > threepm_threshold
    win_given_3pm = float(np.mean(fav_win_mask[mask_3pm])) if np.any(mask_3pm) else fav_base_win

    tov_threshold = int(np.round(np.percentile(fav_tov_arr, 48)))
    mask_tov = fav_tov_arr < tov_threshold
    win_given_tov = float(np.mean(fav_win_mask[mask_tov])) if np.any(mask_tov) else fav_base_win

    ortg_threshold = round(float(np.percentile(fav_ortg_arr, 50)), 1)
    mask_ortg = fav_ortg_arr > ortg_threshold
    win_given_ortg = float(np.mean(fav_win_mask[mask_ortg])) if np.any(mask_ortg) else fav_base_win

    props_list = []
    for side_row, opp_row in [(away_row, home_row), (home_row, away_row)]:
        p_team = str(side_row["Team_Abbr"])
        p_opp = str(opp_row["Team_Abbr"])
        t_pace = float(side_row["Pace"])
        o_def_rtg = float(opp_row["Def_Rtg"])

        for star_idx in (1, 2):
            s_name = str(side_row[f"Star{star_idx}_Name"])
            s_ppg = float(side_row[f"Star{star_idx}_PPG"])
            s_line = float(side_row.get(f"Star{star_idx}_Line", round_to_half_point(s_ppg)))
            s_style = str(side_row.get(f"Star{star_idx}_Style", "ISO")).upper()
            s_std = float(side_row.get(f"Star{star_idx}_Std", 5.6))

            ppg_b_raw = float(side_row.get(f"Star{star_idx}_PPG_Base", s_ppg))
            rpg_b_raw = float(side_row.get(f"Star{star_idx}_RPG_Base", 5.6))
            apg_b_raw = float(side_row.get(f"Star{star_idx}_APG_Base", 4.8))
            min_b_raw = float(side_row.get(f"Star{star_idx}_Minutos_Base", 34.0))
            min_p_raw = float(side_row.get(f"Star{star_idx}_Minutos_Proy", min_b_raw))

            t_style_off = float(side_row["Iso_PPP_Off"]) if s_style == "ISO" else float(side_row["PnR_PPP_Off"])
            o_style_def = float(opp_row["Iso_PPP_Def"]) if s_style == "ISO" else float(opp_row["PnR_PPP_Def"])

            prop_res = simulate_player_prop(
                rng=rng,
                player_name=s_name,
                team_abbr=p_team,
                opp_abbr=p_opp,
                base_ppg=s_ppg,
                prop_line=s_line,
                style=s_style,
                player_std=s_std,
                team_pace=t_pace,
                exp_pace=exp_possessions,
                team_style_off_ppp=t_style_off,
                opp_style_def_ppp=o_style_def,
                opp_def_rtg=o_def_rtg,
                poss_sim=poss_sim,
                ppg_base_raw=ppg_b_raw,
                rpg_base=rpg_b_raw,
                apg_base=apg_b_raw,
                minutos_base=min_b_raw,
                minutos_proyectados=min_p_raw,
            )
            prop_res["matchup_label"] = f"{away_abbr} @ {home_abbr}"
            props_list.append(prop_res)

    return {
        "matchup_id": f"{away_abbr}@{home_abbr}",
        "matchup_title": f"{away_row['Team']} @ {home_row['Team']}",
        "away_team": str(away_row["Team"]),
        "away_abbr": away_abbr,
        "home_team": str(home_row["Team"]),
        "home_abbr": home_abbr,
        "away_row": away_row,
        "home_row": home_row,
        "n_sims": n_sims,
        "exp_possessions": exp_possessions,
        "exp_away_pts": exp_away_pts,
        "exp_home_pts": exp_home_pts,
        "exp_margin_home": exp_margin_home,
        "exp_total_pts": exp_total_pts,
        "proj_spread_home": proj_spread_home,
        "spread_line_home": spread_line_home,
        "spread_line_away": spread_line_away,
        "total_line": total_line,
        "cuota_over_decimal": cuota_over_decimal,
        "cuota_under_decimal": cuota_under_decimal,
        "cuota_spread_home_decimal": cuota_spread_home_decimal,
        "cuota_spread_away_decimal": cuota_spread_away_decimal,
        "odds_source_label": odds_source_label,
        "valor_totales": valor_totales,
        "valor_spread": valor_spread,
        "away_win_prob": away_win_prob,
        "home_win_prob": home_win_prob,
        "away_ml_dec": away_ml_dec,
        "away_ml_amer": away_ml_amer,
        "home_ml_dec": home_ml_dec,
        "home_ml_amer": home_ml_amer,
        "away_cover_prob": away_cover_prob,
        "home_cover_prob": home_cover_prob,
        "away_sp_dec": away_sp_dec,
        "away_sp_amer": away_sp_amer,
        "home_sp_dec": home_sp_dec,
        "home_sp_amer": home_sp_amer,
        "over_prob": over_prob,
        "under_prob": under_prob,
        "over_dec": over_dec,
        "over_amer": over_amer,
        "under_dec": under_dec,
        "under_amer": under_amer,
        "total_p10": float(np.percentile(total_pts_sim, 10)),
        "total_p50": float(np.percentile(total_pts_sim, 50)),
        "total_p90": float(np.percentile(total_pts_sim, 90)),
        "fav_abbr": fav_abbr,
        "fav_name": fav_name,
        "fav_base_win": fav_base_win,
        "threepm_threshold": threepm_threshold,
        "win_given_3pm": win_given_3pm,
        "tov_threshold": tov_threshold,
        "win_given_tov": win_given_tov,
        "ortg_threshold": ortg_threshold,
        "win_given_ortg": win_given_ortg,
        "away_iso_delta": away_iso_delta,
        "away_pnr_delta": away_pnr_delta,
        "home_iso_delta": home_iso_delta,
        "home_pnr_delta": home_pnr_delta,
        "away_tactical_ortg_adj": away_tactical_ortg_adj,
        "home_tactical_ortg_adj": home_tactical_ortg_adj,
        "props": props_list,
    }


def extract_executive_radar(sim_results: List[Dict]) -> Dict:
    """Extrae los 4 Top Picks automáticos en formato estrictamente DECIMAL evaluando los partidos activos."""
    best_ml = None
    best_spread = None
    best_total = None
    best_prop = None

    for res in sim_results:
        ml_candidates = [
            (res["home_team"], res["home_abbr"], res["home_win_prob"], res["home_ml_dec"], f"vs {res['away_abbr']} (Local)", res["exp_home_pts"] - res["exp_away_pts"]),
            (res["away_team"], res["away_abbr"], res["away_win_prob"], res["away_ml_dec"], f"@ {res['home_abbr']} (Visita)", res["exp_away_pts"] - res["exp_home_pts"]),
        ]
        for c in ml_candidates:
            if best_ml is None or c[2] > best_ml["prob"]:
                best_ml = {"team": c[0], "abbr": c[1], "prob": c[2], "dec": c[3], "context": c[4], "proj_margin": c[5]}

        sp_candidates = [
            (f"{res['home_team']} {res['spread_line_home']:+.1f} @ {res['cuota_spread_home_decimal']:.2f}", res["home_abbr"], res["home_cover_prob"], res["home_sp_dec"], f"{res['away_abbr']} @ {res['home_abbr']} | Modelo: {res['proj_spread_home']:+.1f}"),
            (f"{res['away_team']} {res['spread_line_away']:+.1f} @ {res['cuota_spread_away_decimal']:.2f}", res["away_abbr"], res["away_cover_prob"], res["away_sp_dec"], f"{res['away_abbr']} @ {res['home_abbr']} | Modelo: {-res['proj_spread_home']:+.1f}"),
        ]
        for c in sp_candidates:
            if best_spread is None or c[2] > best_spread["prob"]:
                best_spread = {"pick": c[0], "abbr": c[1], "prob": c[2], "dec": c[3], "context": c[4]}

        tot_candidates = [
            (f"OVER {res['total_line']:.1f} @ {res['cuota_over_decimal']:.2f}", res["home_abbr"], res["over_prob"], res["over_dec"], f"{res['away_abbr']} @ {res['home_abbr']} | Modelo: {res['exp_total_pts']:.1f} pts"),
            (f"UNDER {res['total_line']:.1f} @ {res['cuota_under_decimal']:.2f}", res["home_abbr"], res["under_prob"], res["under_dec"], f"{res['away_abbr']} @ {res['home_abbr']} | Modelo: {res['exp_total_pts']:.1f} pts"),
        ]
        for c in tot_candidates:
            if best_total is None or c[2] > best_total["prob"]:
                best_total = {"pick": c[0], "abbr": c[1], "prob": c[2], "dec": c[3], "context": c[4]}

        for prop in res["props"]:
            if best_prop is None or prop["lean_prob"] > best_prop["lean_prob"]:
                best_prop = prop

    return {"lock": best_ml, "spread": best_spread, "total": best_total, "prop": best_prop}